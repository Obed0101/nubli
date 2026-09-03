#!/usr/bin/env python3
"""
Nubli: local-first document anonymizer.

It extracts Markdown/text from common office formats and applies replacement
rules without sending data to any web service. OCR is optional and local.
"""

from __future__ import annotations

import argparse
import csv
from difflib import SequenceMatcher
import hashlib
import importlib
import io
import json
import re
import sys
from time import perf_counter
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Sequence


TEXT_EXTS = {".txt", ".md", ".markdown", ".html", ".htm", ".json", ".xml", ".yml", ".yaml", ".log"}
CSV_EXTS = {".csv", ".tsv"}
DOCX_EXTS = {".docx"}
XLSX_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
PDF_EXTS = {".pdf"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}

TOKEN_RE = re.compile(r"[^\W_]+", re.UNICODE)

DEFAULT_DOCUMENT_LABEL = "DOCUMENT-001"
STRUCTURAL_DELIMITER_PAIRS = {
    "(": ")",
    "[": "]",
    "{": "}",
    "<": ">",
    "“": "”",
    "‘": "’",
    "«": "»",
    "\"": "\"",
    "'": "'",
    "`": "`",
}
STRUCTURAL_OPENERS = set(STRUCTURAL_DELIMITER_PAIRS)
STRUCTURAL_CLOSERS = {value for key, value in STRUCTURAL_DELIMITER_PAIRS.items() if key not in {'"', "'", "`"}}
STRUCTURAL_CHARS = STRUCTURAL_OPENERS | STRUCTURAL_CLOSERS
PLACEHOLDER_RE = re.compile(
    r"^(?:PERSON|COMPANY|EMAIL|PHONE|ID|ADDRESS|ACCOUNT|DATE|LOCATION|VALUE)-\d+(?:@(?:example\.(?:com|org|net)|demo\.local|local))?$",
    re.IGNORECASE,
)
HTML_TAG_NAMES = {
    "a", "abbr", "article", "aside", "b", "body", "br", "button", "caption", "code", "col",
    "div", "em", "figcaption", "figure", "footer", "form", "h1", "h2", "h3", "h4", "h5", "h6",
    "head", "header", "hr", "html", "i", "img", "input", "label", "li", "link", "main", "meta",
    "nav", "ol", "option", "p", "pre", "script", "section", "select", "small", "span", "strong",
    "style", "table", "tbody", "td", "tfoot", "th", "thead", "title", "tr", "u", "ul",
}

PHONE_PATTERN = re.compile(
    r"(?<![\w-])(?:(?:\(\s*\+?[0-9]{1,3}\s*\)\s*)|(?:\+[0-9]{1,3}[\s.-])|(?:507[\s.-]))?"
    r"(?:\(?[0-9]{3}\)?[\s.-][0-9]{4}|\(?[0-9]{4}\)?[\s.-][0-9]{4}|"
    r"\(?[0-9]{3}\)?[\s.-][0-9]{3}[\s.-][0-9]{4}|[0-9]{8})(?![\w-])"
)

PII_PATTERNS: list[tuple[str, re.Pattern[str], str]] = [
    (
        "email",
        re.compile(r"(?<![\w.+-])[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,63}(?![\w.-])", re.IGNORECASE),
        "EMAIL-0001@example.com",
    ),
    (
        "phone",
        PHONE_PATTERN,
        "PHONE-0001",
    ),
    ("id_or_ruc", re.compile(r"(?<!\w)[0-9]{1,3}(?:[-/.\s][0-9]{1,8}){1,3}(?!\w)"), "ID-0001"),
]

ACCOUNT_PATTERN = re.compile(
    r"(?i)(?:\b(?:cuenta(?:\s+bancaria)?|account|iban|swift|bank\s+account|n[úu]mero\s+de\s+cuenta)\b)\s*"
    r"(?:n[úu]m(?:ero)?\.?\s*)?(?:[:#=-]\s*)?"
    r"(?P<value>(?:[A-Z]{2}[0-9]{2}[A-Z0-9]{10,30}|[A-Z0-9]{8,11}|[0-9](?:[0-9 -]{6,24}[0-9])))"
)
DATE_PATTERN = re.compile(
    r"(?<!\w)(?:[0-3]?\d[/-][01]?\d[/-](?:\d{2}|\d{4})|(?:\d{4}[/-][01]?\d[/-][0-3]?\d))(?!\w)"
)
_CONTEXT_LABELS_PATTERN = (
    r"(?:empresa|compa[ñn][íi]a|sociedad|corporaci[óo]n|proveedor|cliente|company|"
    r"organization|organizaci[óo]n|vendor|supplier|bank|banco|direcci[óo]n|domicilio|"
    r"address|residencia|residence|ubicaci[óo]n|lugar|localizaci[óo]n|location|city|"
    r"ciudad|country|pa[íi]s|tel[eé]fono|telephone|phone|tel\.?|id|identificador|"
    r"ruc|pasaporte|passport|tax\s+id|document(?:o|\s+number)?\s*id|cuenta|"
    r"n[úu]mero\s+de\s+cuenta|account|iban|swift|fecha|date|nacimiento|birth|"
    r"nombre|name|representante|representative|apoderado|attorney|valor|value|"
    r"dato|data)"
)
_CONTEXT_END = rf"(?=\s*(?:[,;\[\]\(\)\{{\}}<>«»“”]|{_CONTEXT_LABELS_PATTERN}\s*[:=#-])|$)"
ADDRESS_PATTERN = re.compile(
    r"(?ix)(?<!\w)(?:calle|c\.?|avenida|av\.?|carrera|cra\.?|v[íi]a|boulevard|blvd\.?|street|st\.?|road|rd\.?|avenue|ave\.?)\s+"
    r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9#][A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9#'./-]*"
    rf"(?:[ \t]+(?!{_CONTEXT_LABELS_PATTERN}\s*[:=#-])[A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9#][A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9#'./-]*){{0,10}}"
    r"(?![A-Za-z0-9_.+-]*@)"
)
ORG_PATTERN = re.compile(
    r"(?ix)(?<![\w])(?:[A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9][A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9&.'’-]*[ \t,]*){1,6}(?:S\.?[ \t]+A\.?|S\.A\.?|SAS|L\.?T\.?D\.?|L\.L\.C\.?|I\.?N\.?C\.?|C\.?O\.?R\.?P\.?|GmbH)(?![\w])"
)
COMPANY_CONTEXT_PATTERN = re.compile(
    r"(?ix)(?<![\w])(?P<label>empresa|compa[ñn][íi]a|sociedad|corporaci[óo]n|proveedor|cliente|company|organization|vendor|supplier|bank|banco)\s*[:=-]\s*"
    rf"(?P<value>[^\n\]\)\}}»”>\[]{{2,90}}?){_CONTEXT_END}"
)
LOCATION_CONTEXT_PATTERN = re.compile(
    r"(?ix)(?<![\w])(?:ubicaci[óo]n|lugar|localizaci[óo]n|location|city|ciudad|country|pa[íi]s)\s*[:=-]\s*"
    rf"(?P<value>[^\n\]\)\}}»”>\[]{{2,80}}?){_CONTEXT_END}"
)
ADDRESS_CONTEXT_PATTERN = re.compile(
    r"(?ix)(?<![\w])(?:direcci[óo]n|domicilio|address|residencia|residence)\s*[:=-]\s*"
    rf"(?P<value>[^\n\]\)\}}»”>\[]{{2,100}}?){_CONTEXT_END}"
)
VALUE_CONTEXT_PATTERN = re.compile(
    r"(?ix)(?<![\w])(?:valor|value|dato|data)\s*[:=-]\s*"
    rf"(?P<value>[^\n\]\)\}}»”>\[]{{1,80}}?){_CONTEXT_END}"
)
IDENTIFIER_CONTEXT_PATTERN = re.compile(
    r"(?ix)(?<![\w])(?:id|identificador|ruc|pasaporte|passport|tax\s+id|document(?:o|\s+number)?\s*id)\s*[:#=-]\s*(?P<value>[A-Z0-9][A-Z0-9._\-/]{2,32})"
)
NAME_TOKEN_PATTERN = r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ][A-Za-zÁÉÍÓÚÜÑáéíóúüñ'’\-]{1,30}"
NAME_PATTERN = re.compile(
    rf"(?<![\w]){NAME_TOKEN_PATTERN}(?:(?:[ \t]+|[ \t]*[-/]\s*){NAME_TOKEN_PATTERN}){{1,3}}(?![\w])"
)

# Keep this list boring on purpose. These are common report headings that look
# like person names to a cheap regex; avoiding noisy suggestions matters because
# users will ignore the safety panel if it cries wolf too often.
COMMON_NAME_FALSE_POSITIVES = {
    "Demo Company",
    "Empresa Demo",
    "Persona Demo",
    "Estados Financieros",
    "Estado De",
    "Notas A",
    "República De",
    "Ciudad De",
    "Panamá",
    "Total Activos",
    "Total Pasivos",
}

SAFE_PLACEHOLDER_DOMAINS = (".local", "example.com", "example.org", "example.net")


@dataclass
class Rule:
    source: str
    target: str
    case_sensitive: bool = False
    smart: bool = True
    preserve_case: bool = True
    fuzzy_threshold: float = 0.88


@dataclass
class Suggestion:
    kind: str
    value: str
    target: str
    reason: str
    confidence: float = 0.0
    source: str = "deterministic"


@dataclass(frozen=True)
class SensitiveDetection:
    """A sensitive span in the exact text passed to the detector.

    Positions always refer to the original string.  Delimiters and markup are
    intentionally outside the span, so replacing a value never removes its
    surrounding structure.
    """

    kind: str
    value: str
    start: int
    end: int
    confidence: float
    reason: str
    placeholder_kind: str
    automatic: bool = True
    source: str = "deterministic"


@dataclass(frozen=True)
class StructuralSpan:
    start: int
    end: int
    content_start: int
    content_end: int
    opener: str
    closer: str | None
    depth: int
    balanced: bool


@dataclass(frozen=True)
class ProjectedText:
    visible: str
    positions: tuple[int | None, ...]
    has_markup: bool


@dataclass(frozen=True)
class _Candidate:
    start: int
    end: int
    replacement: str
    kind: str
    index: int = -1
    priority: int = 0


@dataclass(frozen=True)
class _DocxTextSegment:
    text: str
    start: int
    end: int
    bold: bool = False
    italic: bool = False
    underline: bool = False
    highlight: bool = False
    hyperlink: bool = False
    hyperlink_url: str | None = None


@dataclass
class WordBox:
    text: str
    left: float
    top: float
    right: float
    bottom: float
    confidence: float = 100.0


def _is_word_apostrophe(text: str, index: int) -> bool:
    """Return whether a quote is an apostrophe inside a word."""

    if index <= 0 or index + 1 >= len(text):
        return False
    return text[index - 1].isalpha() and text[index + 1].isalpha()


def scan_structural_spans(text: str) -> list[StructuralSpan]:
    """Scan nested delimiters without changing the source string.

    The scanner is deliberately tolerant: valid pairs are returned as
    balanced spans and unmatched opening delimiters become spans extending to
    the end of the input.  An unmatched closing delimiter is left untouched.
    This gives detectors useful context for incomplete documents while keeping
    all replacement offsets anchored to the original text.
    """

    stack: list[tuple[int, str, int]] = []
    spans: list[StructuralSpan] = []

    for index, char in enumerate(text):
        if char in {"'", '"', "`", "‘"} and _is_word_apostrophe(text, index):
            continue

        if char in STRUCTURAL_OPENERS:
            if char in {"'", '"', "`"} and stack and stack[-1][1] == char:
                start, opener, depth = stack.pop()
                spans.append(
                    StructuralSpan(
                        start=start,
                        end=index + 1,
                        content_start=start + 1,
                        content_end=index,
                        opener=opener,
                        closer=char,
                        depth=depth,
                        balanced=True,
                    )
                )
            else:
                stack.append((index, char, len(stack)))
            continue

        if char in STRUCTURAL_CLOSERS:
            if not stack:
                continue
            start, opener, depth = stack[-1]
            expected = STRUCTURAL_DELIMITER_PAIRS[opener]
            if char != expected:
                # Do not close an outer frame through an unmatched nested
                # opener.  The remaining frame will be reported as incomplete.
                continue
            stack.pop()
            spans.append(
                StructuralSpan(
                    start=start,
                    end=index + 1,
                    content_start=start + 1,
                    content_end=index,
                    opener=opener,
                    closer=char,
                    depth=depth,
                    balanced=True,
                )
            )

    for start, opener, depth in stack:
        spans.append(
            StructuralSpan(
                start=start,
                end=len(text),
                content_start=start + 1,
                content_end=len(text),
                opener=opener,
                closer=None,
                depth=depth,
                balanced=False,
            )
        )

    return sorted(spans, key=lambda span: (span.start, span.end, span.depth))


class StructuralScanner:
    """Reusable structural context for one immutable text snapshot."""

    def __init__(self, text: str):
        self.text = text
        self.spans = scan_structural_spans(text)

    def containing(self, start: int, end: int) -> list[StructuralSpan]:
        return [
            span
            for span in self.spans
            if span.content_start <= start and end <= span.content_end
        ]

    def is_inside(self, start: int, end: int) -> bool:
        return bool(self.containing(start, end))

    def crosses_delimiter(self, start: int, end: int) -> bool:
        return any(char in STRUCTURAL_CHARS for char in self.text[start:end])


def structural_spans(text: str) -> list[StructuralSpan]:
    """Compatibility-friendly function name for callers needing span maps."""

    return scan_structural_spans(text)


def _markdown_or_html_projection(text: str, markup: str) -> ProjectedText:
    """Build a visible-text projection while retaining raw character offsets."""

    visible: list[str] = []
    positions: list[int] = []
    index = 0
    has_markup = False
    markup = markup.lower()

    while index < len(text):
        char = text[index]
        if markup == "html" and char == "<":
            close = text.find(">", index + 1)
            if close >= 0 and _looks_like_html_tag(text, index, close + 1):
                has_markup = True
                if visible and not visible[-1].isspace():
                    visible.append(" ")
                    positions.append(None)
                index = close + 1
                continue
        if markup == "markdown":
            if char in "*`":
                has_markup = True
                index += 1
                continue
            if char in "[]()" and not (index > 0 and text[index - 1] == "\\"):
                has_markup = True
                if char == "]" and index + 1 < len(text) and text[index + 1] == "(":
                    close = text.find(")", index + 2)
                    if close >= 0:
                        index = close + 1
                        continue
                index += 1
                continue
            if char == "_":
                previous = text[index - 1] if index else " "
                following = text[index + 1] if index + 1 < len(text) else " "
                if not (previous.isalnum() and following.isalnum()):
                    has_markup = True
                    index += 1
                    continue
        visible.append(char)
        positions.append(index)
        index += 1

    return ProjectedText("".join(visible), tuple(positions), has_markup)


def die(message: str, code: int = 2) -> None:
    print(f"nubli: {message}", file=sys.stderr)
    raise SystemExit(code)


def require_import(module_name: str, install_hint: str):
    try:
        return importlib.import_module(module_name)
    except ImportError as exc:
        die(f"missing dependency '{module_name}'. Install it with: {install_hint}")
        raise exc


def normalize_token(value: str, case_sensitive: bool = False) -> str:
    cleaned = re.sub(r"[\W_]+", "", value, flags=re.UNICODE)
    return cleaned if case_sensitive else cleaned.casefold()


def source_tokens(value: str) -> list[str]:
    return [token for token in TOKEN_RE.findall(value) if token]


def _looks_like_html_tag(text: str, start: int, end: int) -> bool:
    if start < 0 or end <= start or end > len(text) or text[start] != "<" or text[end - 1] != ">":
        return False
    inner = text[start + 1:end - 1].strip()
    if inner.casefold().startswith(("!--", "!doctype", "![cdata[", "?")):
        return True
    match = re.fullmatch(r"/?([A-Za-z][A-Za-z0-9:-]*)(?:\s+[^>]*)?/?", inner)
    if not match:
        return False
    tag_name = match.group(1).casefold()
    return tag_name in HTML_TAG_NAMES or bool(re.search(r"\s+[A-Za-z_:][\w:.-]*\s*=", inner))


def similarity(left: str, right: str) -> float:
    if not left or not right:
        return 0.0
    return SequenceMatcher(None, left, right).ratio()


def prompt_yes_no(message: str, default: bool = False) -> bool:
    marker = "Y/n" if default else "y/N"
    answer = input(f"{message} [{marker}] ").strip().lower()
    if not answer:
        return default
    return answer in {"y", "yes", "s", "si", "sí"}


def redact_preview(value: str, keep: int = 3) -> str:
    value = value.strip()
    if len(value) <= keep * 2:
        return "*" * len(value)
    return f"{value[:keep]}…{value[-keep:]}"


def safe_placeholder(kind: str, value: str) -> bool:
    lowered = value.strip().casefold()
    if PLACEHOLDER_RE.fullmatch(value.strip()):
        return True
    if kind.casefold() == "email" and "@" in lowered:
        domain = lowered.rsplit("@", 1)[1]
        return domain.endswith(SAFE_PLACEHOLDER_DOMAINS)
    return lowered in {
        "tel-0000",
        "id-0000",
        "persona demo",
        "empresa demo",
        "demo company",
        "person demo",
        "company demo",
    }


GENERIC_NAME_WORDS = {
    "a",
    "an",
    "and",
    "company",
    "confidential",
    "document",
    "ejemplo",
    "example",
    "financieros",
    "general",
    "heading",
    "interna",
    "internal",
    "legal",
    "nota",
    "note",
    "normal",
    "sample",
    "text",
    "texto",
    "the",
    "total",
    "variant",
    "uppercase",
    "lowercase",
    "bold",
    "italic",
    "underlined",
    "table",
    "list",
    "valor",
    "value",
    "dato",
    "data",
    "location",
    "city",
    "ciudad",
    "country",
    "país",
}
NAME_CONTEXT_WORDS = {
    "cliente",
    "client",
    "contact",
    "contacto",
    "customer",
    "name",
    "nombre",
    "person",
    "persona",
    "representative",
    "representante",
    "apoderado",
    "attorney",
}
NON_PERSON_ENTITY_WORDS = {
    "account",
    "address",
    "avenida",
    "bancaria",
    "bank",
    "calle",
    "city",
    "ciudad",
    "company",
    "compañía",
    "cuenta",
    "dirección",
    "domicilio",
    "empresa",
    "id",
    "iban",
    "location",
    "organización",
    "organization",
    "pasaporte",
    "passport",
    "phone",
    "ruc",
    "teléfono",
    "telephone",
    "ubicación",
    "valor",
    "value",
    "dato",
    "data",
}
ORG_CONTEXT_WORDS = {"company", "compañía", "empresa", "organization", "organización", "org"}
DETECTION_PRIORITIES = {
    "email": 120,
    "phone": 100,
    "id_or_ruc": 115,
    "account": 95,
    "address": 90,
    "organization": 85,
    "date": 80,
    "location": 78,
    "person": 75,
    "possible_person": 70,
    "value": 60,
}
IDENTIFIER_LABEL_WORDS = {
    "id",
    "identificador",
    "ruc",
    "pasaporte",
    "passport",
    "tax",
    "documento",
    "document",
}


def canonical_sensitive_value(kind: str, value: str) -> str:
    value = value.strip()
    if kind in {"email", "organization", "person", "possible_person", "address"}:
        return re.sub(r"\s+", " ", value).casefold()
    return re.sub(r"[^\w]+", "", value, flags=re.UNICODE).casefold()


def placeholder_kind_for(kind: str) -> str:
    return {
        "email": "EMAIL",
        "phone": "PHONE",
        "id_or_ruc": "ID",
        "organization": "COMPANY",
        "person": "PERSON",
        "possible_person": "PERSON",
        "address": "ADDRESS",
        "account": "ACCOUNT",
        "date": "DATE",
    }.get(kind, kind.upper())


def default_placeholder(kind: str) -> str:
    prefix = placeholder_kind_for(kind)
    return f"{prefix}-0001@example.com" if prefix == "EMAIL" else f"{prefix}-0001"


def _context_has_label(text: str, start: int, labels: set[str]) -> bool:
    before = text[max(0, start - 48):start].casefold()
    words = set(re.findall(r"[a-záéíóúüñ]+", before, flags=re.IGNORECASE))
    return bool(words & labels)


def _name_detection_is_plausible(
    text: str,
    match: re.Match[str],
    scanner: StructuralScanner,
) -> tuple[bool, float, bool]:
    return _name_value_is_plausible(text, match.group(0), match.start(), match.end(), scanner)


def _name_value_is_plausible(
    text: str,
    raw_value: str,
    start: int,
    end: int,
    scanner: StructuralScanner,
) -> tuple[bool, float, bool]:
    value = raw_value.strip(" \t,;:.\"'“”«»`()[]{}")
    tokens = source_tokens(value)
    if not 2 <= len(tokens) <= 4:
        return False, 0.0, False
    lowered = {token.casefold() for token in tokens}
    if value.casefold() in {item.casefold() for item in COMMON_NAME_FALSE_POSITIVES}:
        return False, 0.0, False
    if lowered <= GENERIC_NAME_WORDS:
        return False, 0.0, False
    if lowered <= NAME_CONTEXT_WORDS:
        return False, 0.0, False
    if lowered & NON_PERSON_ENTITY_WORDS:
        return False, 0.0, False

    title = all(token[:1].isupper() and token[1:].islower() for token in tokens if len(token) > 1)
    uppercase = all(token.upper() == token for token in tokens if any(ch.isalpha() for ch in token))
    bounded = scanner.is_inside(start, end)
    labeled = _context_has_label(text, start, NAME_CONTEXT_WORDS)
    has_name_word = bool(lowered & NAME_CONTEXT_WORDS)
    if lowered & GENERIC_NAME_WORDS and not (bounded or labeled or has_name_word):
        return False, 0.0, False

    # Lowercase phrases need a person-specific label.  Delimiters are useful
    # boundaries, but they do not turn arbitrary quoted values into names.
    if not (title or uppercase or labeled or has_name_word):
        return False, 0.0, False
    confidence = 0.96 if bounded or labeled else 0.91 if title or uppercase else 0.88
    automatic = labeled or title or uppercase
    return True, confidence, automatic


def _iter_name_candidates(text: str) -> Iterable[tuple[int, int, str]]:
    token_matches = list(re.finditer(NAME_TOKEN_PATTERN, text))
    for start_index, first in enumerate(token_matches):
        for count in range(2, 5):
            end_index = start_index + count - 1
            if end_index >= len(token_matches):
                break
            valid = True
            for index in range(start_index, end_index):
                separator = text[token_matches[index].end():token_matches[index + 1].start()]
                if not re.fullmatch(r"(?:[ \t]+|[ \t]*[-/]\s*)", separator):
                    valid = False
                    break
            if not valid:
                break
            last = token_matches[end_index]
            yield first.start(), last.end(), text[first.start():last.end()]


def _clean_context_value(value: str) -> str:
    value = re.split(r"(?i)(?=https?://)", value, maxsplit=1)[0]
    value = re.split(r"(?i)(?=\b[\w.+-]+@[\w.-]+\.)", value, maxsplit=1)[0]
    value = value.strip(" \t\r\n\"'“”«»`()[]{}<>*_~")
    return value.rstrip(" \t,;:.")


def _cleaned_context_span(text: str, start: int, end: int) -> tuple[str, int, int]:
    raw = text[start:end]
    value = _clean_context_value(raw)
    if not value:
        return "", start, start
    relative = raw.find(value)
    if relative < 0:
        relative = len(raw) - len(raw.lstrip(" \t\r\n\"'“”«»`()[]{}<>"))
    value_start = start + relative
    return value, value_start, value_start + len(value)


def _html_tag_context(text: str, start: int, end: int) -> tuple[bool, bool]:
    """Return (inside_tag, relevant_attribute) for one candidate span."""

    opening = text.rfind("<", 0, start)
    closing = text.rfind(">", 0, start)
    tag_end = text.find(">", opening + 1) if opening >= 0 else -1
    if opening <= closing or tag_end < 0 or tag_end < end or not _looks_like_html_tag(text, opening, tag_end + 1):
        return False, False
    tag_prefix = text[opening:start].casefold()
    relevant = bool(re.search(r"(?:alt|title|aria-label|href|src|name|value)\s*=\s*[^=]*$", tag_prefix))
    return True, relevant


def _is_placeholder_label(text: str, start: int, value_start: int) -> bool:
    prefix = text[start:value_start].strip()
    return bool(
        re.fullmatch(
            r"(?i)(?:PERSON|COMPANY|EMAIL|PHONE|ID|ADDRESS|ACCOUNT|DATE|LOCATION|VALUE)-",
            prefix,
        )
    )


def detect_sensitive_content(text: str, include_low_confidence: bool = True) -> list[SensitiveDetection]:
    """Return deterministic sensitive spans using structural context.

    This function performs detection only.  Replacement policy is applied by
    :class:`ReplacementEngine`, allowing callers to inspect low-confidence
    suggestions without silently rewriting them.
    """

    scanner = StructuralScanner(text)
    detections: list[SensitiveDetection] = []

    def add(
        kind: str,
        value: str,
        start: int,
        end: int,
        confidence: float,
        reason: str,
        automatic: bool = True,
        placeholder_kind: str | None = None,
        source: str = "deterministic",
    ) -> None:
        value = value.strip()
        if not value or safe_placeholder(kind, value):
            return
        detections.append(
            SensitiveDetection(
                kind=kind,
                value=value,
                start=start,
                end=end,
                confidence=confidence,
                reason=reason,
                placeholder_kind=placeholder_kind or placeholder_kind_for(kind),
                automatic=automatic,
                source=source,
            )
        )

    for kind, pattern, _target in PII_PATTERNS:
        for match in pattern.finditer(text):
            if kind == "phone" and _context_has_label(text, match.start(), IDENTIFIER_LABEL_WORDS):
                continue
            if kind == "id_or_ruc" and (
                DATE_PATTERN.fullmatch(match.group(0).strip())
                or PHONE_PATTERN.fullmatch(match.group(0).strip())
            ):
                continue
            add(kind, match.group(0), match.start(), match.end(), 0.99, "deterministic pattern")

    for match in IDENTIFIER_CONTEXT_PATTERN.finditer(text):
        if _is_placeholder_label(text, match.start(), match.start("value")):
            continue
        value, value_start, value_end = _cleaned_context_span(text, match.start("value"), match.end("value"))
        if value:
            add(
                "id_or_ruc",
                value,
                value_start,
                value_end,
                0.99,
                "identifier label",
            )

    for match in ACCOUNT_PATTERN.finditer(text):
        if _is_placeholder_label(text, match.start(), match.start("value")):
            continue
        value = match.group("value").strip()
        compact_value = re.sub(r"[\s-]", "", value)
        if len(compact_value) >= 8:
            add("account", value, match.start("value"), match.end("value"), 0.96, "account label")

    for match in DATE_PATTERN.finditer(text):
        context = text[max(0, match.start() - 48):match.start()].casefold()
        associated = bool(re.search(r"(?:nacimiento|birth|dob|fecha\s+de\s+nac|born|fecha|date)", context))
        add(
            "date",
            match.group(0),
            match.start(),
            match.end(),
            0.93 if associated else 0.90,
            "identity-associated date" if associated else "date pattern",
        )

    for match in ADDRESS_PATTERN.finditer(text):
        value, value_start, value_end = _cleaned_context_span(text, match.start(), match.end())
        if len(source_tokens(value)) >= 2:
            add("address", value, value_start, value_end, 0.91, "address pattern")

    for match in ADDRESS_CONTEXT_PATTERN.finditer(text):
        if _is_placeholder_label(text, match.start(), match.start("value")):
            continue
        value, value_start, value_end = _cleaned_context_span(text, match.start("value"), match.end("value"))
        if 2 <= len(source_tokens(value)) <= 12:
            add(
                "address",
                value,
                value_start,
                value_end,
                0.97,
                "address label",
            )

    for match in LOCATION_CONTEXT_PATTERN.finditer(text):
        if _is_placeholder_label(text, match.start(), match.start("value")):
            continue
        value, value_start, value_end = _cleaned_context_span(text, match.start("value"), match.end("value"))
        if 1 <= len(source_tokens(value)) <= 8:
            add(
                "location",
                value,
                value_start,
                value_end,
                0.94,
                "location label",
            )

    for match in VALUE_CONTEXT_PATTERN.finditer(text):
        if _is_placeholder_label(text, match.start(), match.start("value")):
            continue
        value, value_start, value_end = _cleaned_context_span(text, match.start("value"), match.end("value"))
        if 1 <= len(source_tokens(value)) <= 10:
            add(
                "value",
                value,
                value_start,
                value_end,
                0.88,
                "value label",
            )

    for match in COMPANY_CONTEXT_PATTERN.finditer(text):
        if _is_placeholder_label(text, match.start(), match.start("value")):
            continue
        value, value_start, value_end = _cleaned_context_span(text, match.start("value"), match.end("value"))
        label = match.group("label").casefold()
        corporate_signal = bool(
            re.search(
                r"(?i)(?:s\.?\s*a\.?|s\.?\s*a\.?s\.?|ltd|llc|inc|corp|empresa|compa[ñn][íi]a|sociedad|organizaci[óo]n|organization|banco|bank)",
                value,
            )
        )
        if label in {"cliente", "client"} and not corporate_signal:
            continue
        if 1 <= len(source_tokens(value)) <= 10:
            add(
                "organization",
                value,
                value_start,
                value_end,
                0.95,
                "organization label",
            )

    for match in ORG_PATTERN.finditer(text):
        raw_value = match.group(0)
        value = _clean_context_value(raw_value)
        for marker in ("compañía", "compania", "empresa", "sociedad", "corporación", "organización", "company", "organization"):
            marker_match = re.search(rf"(?i)\b{re.escape(marker)}\b", value)
            if marker_match:
                value = value[marker_match.start():]
                break
        inside_tag, relevant_attribute = _html_tag_context(text, match.start(), match.end())
        if (not inside_tag or relevant_attribute) and value.casefold() not in {item.casefold() for item in COMMON_NAME_FALSE_POSITIVES}:
            raw_offset = raw_value.find(value)
            raw_offset = max(0, raw_offset)
            add("organization", value, match.start() + raw_offset, match.start() + raw_offset + len(value), 0.94, "organization suffix")

    for start, end, raw_value in _iter_name_candidates(text):
        inside_tag, relevant_attribute = _html_tag_context(text, start, end)
        if inside_tag and not relevant_attribute:
            continue
        plausible, confidence, automatic = _name_value_is_plausible(text, raw_value, start, end, scanner)
        if plausible and (include_low_confidence or automatic):
            value = raw_value.strip(" \t,;:.\"'“”«»`()[]{}")
            add(
                "person" if automatic else "possible_person",
                value,
                start,
                start + len(value),
                confidence,
                "structured name-like phrase" if scanner.is_inside(start, end) else "name-like phrase",
                automatic=automatic,
            )

    # Prefer the most specific/largest span at one location.  The replacement
    # engine performs the final cross-kind overlap decision and records it.
    unique: dict[tuple[str, int, int], SensitiveDetection] = {}
    for detection in detections:
        key = (detection.kind, detection.start, detection.end)
        previous = unique.get(key)
        if previous is None or detection.confidence > previous.confidence:
            unique[key] = detection
    return sorted(
        unique.values(),
        key=lambda item: (item.start, -(item.end - item.start), -DETECTION_PRIORITIES.get(item.kind, 0)),
    )


def suggestion_from_detection(detection: SensitiveDetection, target: str | None = None) -> Suggestion:
    return Suggestion(
        kind=detection.kind,
        value=detection.value,
        target=target or default_placeholder(detection.placeholder_kind),
        reason=detection.reason,
        confidence=detection.confidence,
        source=detection.source,
    )


def merge_suggestions(*groups: Sequence[Suggestion], max_items: int = 80) -> list[Suggestion]:
    merged: list[Suggestion] = []
    seen: set[tuple[str, str]] = set()
    for group in groups:
        for suggestion in group:
            key = (suggestion.kind, canonical_sensitive_value(suggestion.kind, suggestion.value))
            if key in seen:
                continue
            seen.add(key)
            merged.append(suggestion)
            if len(merged) >= max_items:
                return merged
    return merged


def suggest_sensitive_content(text: str, max_items: int = 80) -> list[Suggestion]:
    detections = detect_sensitive_content(text, include_low_confidence=True)
    return merge_suggestions(
        [suggestion_from_detection(detection) for detection in detections],
        max_items=max_items,
    )


def print_suggestions(suggestions: list[Suggestion]) -> None:
    if not suggestions:
        return
    print("\nNubli found possible sensitive values still present:", file=sys.stderr)
    for index, suggestion in enumerate(suggestions[:20], start=1):
        print(
            f"  - {suggestion.kind} #{index}: suggest '{suggestion.target}' "
            f"(confidence={suggestion.confidence:.2f}; {suggestion.reason})",
            file=sys.stderr,
        )
    if len(suggestions) > 20:
        print(f"  ... and {len(suggestions) - 20} more suggestions", file=sys.stderr)


def cleanup_paths(paths: Iterable[Path]) -> None:
    for path in paths:
        try:
            path.unlink()
        except FileNotFoundError:
            pass


def enforce_safety(args: argparse.Namespace, engine: "ReplacementEngine", suggestions: list[Suggestion], written: Iterable[Path] = ()) -> None:
    total_matches = sum(engine.counts) + sum(engine.auto_counts.values()) + engine.visual_redactions
    problems: list[str] = []
    if total_matches == 0 and not args.allow_zero_replacements:
        problems.append("zero replacements/redactions matched")
    if args.require_all_rules:
        missing = [
            f"rule-{index + 1:03d}"
            for index, count in enumerate(engine.effective_rule_counts())
            if count == 0
        ]
        if missing:
            problems.append("rules with zero matches: " + ", ".join(missing[:5]))

    if suggestions:
        print_suggestions(suggestions)
        if args.strict_pii:
            problems.append(f"{len(suggestions)} possible sensitive values still present")

    if not problems:
        return

    message = "Nubli safety check: " + "; ".join(problems)
    if args.interactive and sys.stdin.isatty():
        print(f"\n{message}", file=sys.stderr)
        if prompt_yes_no("Keep/write output anyway?", default=False):
            print("nubli: continuing because user confirmed interactively", file=sys.stderr)
            return

    # Fail closed. A false negative here can leak a client name; a false
    # positive only asks the user to be explicit.
    cleanup_paths(written)
    die(message + "; refused to write output")


def title_like(value: str) -> bool:
    words = [w for w in source_tokens(value) if any(ch.isalpha() for ch in w)]
    return bool(words) and all(w[:1].isupper() and w[1:].islower() for w in words if len(w) > 1)


def custom_case(value: str) -> bool:
    letters = [ch for ch in value if ch.isalpha()]
    if not letters:
        return False
    return not (value.upper() == value or value.lower() == value or title_like(value))


def adapt_case(original: str, replacement: str) -> str:
    has_alpha = any(ch.isalpha() for ch in original)
    if not has_alpha:
        return replacement
    if custom_case(replacement) and original.upper() != original:
        return replacement
    if original.upper() == original:
        return replacement.upper()
    if original.lower() == original:
        return replacement.lower()
    if title_like(original):
        return replacement.title() if replacement.lower() == replacement else replacement
    return replacement


def build_pattern(source: str, smart: bool) -> str:
    if not smart:
        return re.escape(source)

    tokens = source_tokens(source)
    if not tokens:
        return re.escape(source)

    # Allows punctuation/space variants without ever crossing a structural
    # delimiter.  Crossing `) (` would otherwise consume delimiters and could
    # destroy two independent entities in one replacement.
    sep = r"[\s.,;:/\\\-–—_]*"
    body = sep.join(re.escape(token) for token in tokens)
    return rf"(?<![\w]){body}(?![\w])"


def _apply_candidates(text: str, candidates: Sequence[_Candidate]) -> str:
    for candidate in sorted(candidates, key=lambda item: (item.start, item.end), reverse=True):
        text = text[:candidate.start] + candidate.replacement + text[candidate.end:]
    return text


def _candidate_original_span(candidate: _Candidate, origin_map: Sequence[tuple[int, ...]]) -> tuple[int, int] | None:
    affected = [position for item in origin_map[candidate.start:candidate.end] for position in item]
    if not affected:
        return None
    return min(affected), max(affected) + 1


def _advance_origin_map(origin_map: list[tuple[int, ...]], candidates: Sequence[_Candidate]) -> list[tuple[int, ...]]:
    updated = list(origin_map)
    for candidate in sorted(candidates, key=lambda item: item.start, reverse=True):
        affected = tuple(position for item in updated[candidate.start:candidate.end] for position in item)
        replacement_map = [affected] * len(candidate.replacement)
        updated[candidate.start:candidate.end] = replacement_map
    return updated


def _replace_projected_characters(text: str, positions: Sequence[int | None], candidates: Sequence[_Candidate]) -> str:
    """Apply visible-text edits while preserving skipped markup characters."""

    chars = list(text)
    for candidate in sorted(candidates, key=lambda item: item.start, reverse=True):
        if candidate.start < 0 or candidate.end > len(positions) or candidate.start >= candidate.end:
            continue
        raw_positions = [position for position in positions[candidate.start:candidate.end] if position is not None]
        if not raw_positions:
            continue
        first = raw_positions[0]
        selected = set(raw_positions)
        chars[first] = candidate.replacement
        for raw_index in selected - {first}:
            chars[raw_index] = ""
    return "".join(chars)


class LocalNER:
    """Lazy, offline-only adapter for an optional spaCy model."""

    LABELS = {
        "PERSON": "person",
        "PER": "person",
        "ORG": "organization",
        "COMPANY": "organization",
        "LOC": "location",
        "GPE": "location",
        "LOCATION": "location",
    }

    def __init__(self, enabled: bool = False, model_path: str | None = None, language: str = "es", threshold: float = 0.85, suggest_only: bool = False):
        self.enabled = enabled
        self.model_path = model_path
        self.language = language
        self.threshold = threshold
        self.suggest_only = suggest_only
        self._attempted = False
        self.nlp: Any | None = None
        self.error: str | None = None
        self.load_seconds = 0.0
        self.memory_before_bytes = 0
        self.memory_after_bytes = 0
        self.model_disk_bytes = 0

    @staticmethod
    def _rss_bytes() -> int:
        try:
            import resource

            value = int(resource.getrusage(resource.RUSAGE_SELF).ru_maxrss)
            return value if sys.platform == "darwin" else value * 1024
        except (ImportError, OSError, ValueError):
            return 0

    @staticmethod
    def _directory_size(path: Path) -> int:
        if not path.exists():
            return 0
        if path.is_file():
            try:
                return path.stat().st_size
            except OSError:
                return 0
        total = 0
        for child in path.rglob("*"):
            try:
                if child.is_file():
                    total += child.stat().st_size
            except OSError:
                continue
        return total

    def load(self) -> Any | None:
        if self._attempted:
            return self.nlp
        self._attempted = True
        if not self.enabled:
            return None

        if self.model_path and self.model_path.startswith(("http://", "https://")):
            self.error = "remote model paths are disabled"
            return None

        self.memory_before_bytes = self._rss_bytes()
        started = perf_counter()
        try:
            spacy = importlib.import_module("spacy")
            model_ref = str(Path(self.model_path).expanduser()) if self.model_path else f"{self.language}_core_news_sm"
            if self.model_path and not Path(self.model_path).expanduser().exists():
                self.error = "local NER model path does not exist"
                return None
            # spacy.load loads an installed package/path only.  It never calls
            # a downloader; missing models fall back to deterministic rules.
            self.nlp = spacy.load(model_ref)
            model_path = Path(self.model_path).expanduser() if self.model_path else None
            if model_path:
                self.model_disk_bytes = self._directory_size(model_path)
            else:
                try:
                    spec = importlib.util.find_spec(model_ref)
                    if spec and spec.submodule_search_locations:
                        self.model_disk_bytes = self._directory_size(Path(next(iter(spec.submodule_search_locations))))
                except (ImportError, OSError, StopIteration):
                    pass
        except (ImportError, OSError, RuntimeError, ValueError) as exc:
            # Keep the message technical and never include document text.
            self.error = type(exc).__name__
            self.nlp = None
        finally:
            self.load_seconds = perf_counter() - started
            self.memory_after_bytes = self._rss_bytes()
        return self.nlp

    def detect(self, text: str) -> list[SensitiveDetection]:
        model = self.load()
        if model is None:
            return []
        detections: list[SensitiveDetection] = []
        try:
            doc = model(text)
        except (OSError, RuntimeError, ValueError):
            self.error = "model inference failed"
            return []
        for entity in getattr(doc, "ents", ()):
            value = str(getattr(entity, "text", "")).strip()
            start = int(getattr(entity, "start_char", -1))
            end = int(getattr(entity, "end_char", -1))
            label = str(getattr(entity, "label_", "")).upper()
            if not value or start < 0 or end <= start or safe_placeholder(label, value):
                continue
            kind = self.LABELS.get(label, "value")
            score = getattr(entity, "score", None)
            if score is None:
                try:
                    score = float(entity._.confidence)  # type: ignore[attr-defined]
                except (AttributeError, TypeError, ValueError):
                    # Standard spaCy entities do not expose calibrated scores.
                    # Treating an unknown score as high confidence would make
                    # optional NER violate the fail-closed replacement policy.
                    score = 0.0
            score = max(0.0, min(1.0, float(score)))
            detections.append(
                SensitiveDetection(
                    kind=kind,
                    value=value,
                    start=start,
                    end=end,
                    confidence=score,
                    reason=f"local NER {label or 'unknown'}",
                    placeholder_kind=placeholder_kind_for(kind),
                    automatic=True,
                    source="spacy",
                )
            )
        return detections

    def report(self) -> dict[str, object]:
        return {
            "enabled": self.enabled,
            "available": self.nlp is not None,
            "language": self.language,
            "threshold": self.threshold,
            "suggest_only": self.suggest_only,
            "load_seconds": round(self.load_seconds, 6),
            "memory_before_bytes": self.memory_before_bytes,
            "memory_after_bytes": self.memory_after_bytes,
            "model_disk_bytes": self.model_disk_bytes,
            "fallback": self.error is not None or (self.enabled and self.nlp is None),
            "error": self.error,
        }


class ReplacementEngine:
    def __init__(self, rules: Iterable[Rule], ner: LocalNER | None = None, auto_pii: bool = True):
        self.rules = list(rules)
        self.counts = [0 for _ in self.rules]
        self.overlap_absorbed = [0 for _ in self.rules]
        self.overlap_events: list[dict[str, object]] = []
        self.auto_counts: dict[str, int] = {}
        self.visual_redactions = 0
        self.visual_review_required = False
        self.ner = ner
        self.auto_pii = auto_pii
        self.suggestions: list[Suggestion] = []
        self._placeholder_by_value: dict[str, str] = {}
        self._placeholder_kind_by_value: dict[str, str] = {}
        self.compiled: list[tuple[Rule, re.Pattern[str]]] = []

        for rule in self.rules:
            flags = 0 if rule.case_sensitive else re.IGNORECASE
            self.compiled.append((rule, re.compile(build_pattern(rule.source, rule.smart), flags)))

    def _protected_values(self) -> set[str]:
        return {
            canonical_sensitive_value("value", rule.target)
            for rule in self.rules
            if rule.target
        }

    def _is_protected(self, kind: str, value: str) -> bool:
        if safe_placeholder(kind, value):
            return True
        normalized = canonical_sensitive_value("value", value)
        return normalized in self._protected_values()

    def _fuzzy_candidates(self, text: str, rule_index: int, rule: Rule) -> list[_Candidate]:
        if rule.fuzzy_threshold <= 0:
            return []
        target = "".join(normalize_token(token, case_sensitive=rule.case_sensitive) for token in source_tokens(rule.source))
        if len(target) < 5:
            return []
        replacement_norm = "".join(
            normalize_token(token, case_sensitive=rule.case_sensitive) for token in source_tokens(rule.target)
        )
        token_matches = list(TOKEN_RE.finditer(text))
        candidates: list[_Candidate] = []
        max_window = min(8, max(1, len(source_tokens(rule.source)) + 3))
        scanner = StructuralScanner(text)

        for start in range(len(token_matches)):
            collected = ""
            for end in range(start, min(len(token_matches), start + max_window)):
                collected += normalize_token(token_matches[end].group(0), case_sensitive=rule.case_sensitive)
                if len(collected) > max(len(target) + 6, int(len(target) * 1.45)):
                    break
                if collected == replacement_norm:
                    continue
                span = (token_matches[start].start(), token_matches[end].end())
                if scanner.crosses_delimiter(*span):
                    continue
                if similarity(collected, target) >= rule.fuzzy_threshold:
                    original = text[span[0]:span[1]]
                    value = adapt_case(original, rule.target) if rule.preserve_case else rule.target
                    candidates.append(
                        _Candidate(
                            start=span[0],
                            end=span[1],
                            replacement=value,
                            kind="rule",
                            index=rule_index,
                            priority=850,
                        )
                    )
                    break
        return candidates

    def _select_candidates(self, candidates: Sequence[_Candidate]) -> list[_Candidate]:
        unique: dict[tuple[int, int, int, str], _Candidate] = {}
        for candidate in candidates:
            key = (candidate.start, candidate.end, candidate.index, candidate.kind)
            previous = unique.get(key)
            if previous is None or candidate.priority > previous.priority:
                unique[key] = candidate

        selected: list[_Candidate] = []
        for candidate in sorted(
            unique.values(),
            key=lambda item: (-item.priority, -(item.end - item.start), item.start, item.index),
        ):
            winner = next(
                (
                    old
                    for old in selected
                    if not (candidate.end <= old.start or candidate.start >= old.end)
                ),
                None,
            )
            if winner is not None:
                if candidate.kind == "rule" and candidate.index >= 0:
                    self.overlap_absorbed[candidate.index] += 1
                    if winner.kind == "rule" and winner.index >= 0:
                        winner_id = f"rule-{winner.index + 1:04d}"
                    else:
                        winner_id = winner.kind
                    self.overlap_events.append(
                        {
                            "absorbed_rule": f"rule-{candidate.index + 1:04d}",
                            "winner": winner_id,
                            "reason": "overlapping longer span",
                        }
                    )
                continue
            selected.append(candidate)
        return sorted(selected, key=lambda item: (item.start, item.end))

    def _apply_rule_stage(self, text: str) -> tuple[str, list[_Candidate]]:
        candidates: list[_Candidate] = []
        for index, (rule, regex) in enumerate(self.compiled):
            for match in regex.finditer(text):
                if match.start() == match.end():
                    continue
                replacement = adapt_case(match.group(0), rule.target) if rule.preserve_case else rule.target
                candidates.append(
                    _Candidate(
                        start=match.start(),
                        end=match.end(),
                        replacement=replacement,
                        kind="rule",
                        index=index,
                        priority=1000,
                    )
                )
            candidates.extend(self._fuzzy_candidates(text, index, rule))

        selected = self._select_candidates(candidates)
        for candidate in selected:
            if candidate.kind == "rule" and candidate.index >= 0:
                self.counts[candidate.index] += 1
        return _apply_candidates(text, selected), selected

    def _record_suggestion(self, suggestion: Suggestion) -> None:
        key = (suggestion.kind, canonical_sensitive_value(suggestion.kind, suggestion.value))
        if any((item.kind, canonical_sensitive_value(item.kind, item.value)) == key for item in self.suggestions):
            return
        self.suggestions.append(suggestion)

    def placeholder_for(self, kind: str, value: str) -> str:
        normalized = canonical_sensitive_value("value", value)
        existing = self._placeholder_by_value.get(normalized)
        if existing:
            return existing
        prefix = placeholder_kind_for(kind)
        self._placeholder_kind_by_value[normalized] = prefix
        used = {
            int(match.group(1))
            for placeholder in self._placeholder_by_value.values()
            if (match := re.search(r"-(\d{4})", placeholder))
            and placeholder.startswith(prefix + "-")
        }
        number = 1
        while number in used:
            number += 1
        placeholder = f"{prefix}-{number:04d}"
        if prefix == "EMAIL":
            placeholder += "@example.com"
        self._placeholder_by_value[normalized] = placeholder
        return placeholder

    def _candidate_for_detection(self, detection: SensitiveDetection) -> _Candidate:
        placeholder = self.placeholder_for(detection.placeholder_kind, detection.value)
        return _Candidate(
            start=detection.start,
            end=detection.end,
            replacement=placeholder,
            kind=detection.kind,
            index=-1,
            priority=DETECTION_PRIORITIES.get(detection.kind, 50),
        )

    def _register_auto_match(self, kind: str) -> None:
        self.auto_counts[kind] = self.auto_counts.get(kind, 0) + 1

    def replacement_for_detection(self, detection: SensitiveDetection) -> str:
        replacement = self.placeholder_for(detection.placeholder_kind, detection.value)
        self._register_auto_match(detection.kind)
        return replacement

    def _apply_deterministic_stage(self, text: str) -> tuple[str, list[_Candidate]]:
        detections = detect_sensitive_content(text, include_low_confidence=True)
        candidates: list[_Candidate] = []
        pending: list[SensitiveDetection] = []
        for detection in detections:
            if not detection.automatic or self._is_protected(detection.kind, detection.value):
                pending.append(detection)
            else:
                candidates.append(self._candidate_for_detection(detection))

        selected = self._select_candidates(candidates)
        selected_ranges = [(candidate.start, candidate.end) for candidate in selected]
        for detection in pending:
            if not any(not (detection.end <= start or detection.start >= end) for start, end in selected_ranges):
                self._record_suggestion(suggestion_from_detection(detection, self.placeholder_for(detection.placeholder_kind, detection.value)))
        for candidate in selected:
            self._register_auto_match(candidate.kind)
        return _apply_candidates(text, selected), selected

    def _apply_ner_stage(self, text: str) -> tuple[str, list[_Candidate]]:
        if self.ner is None:
            return text, []
        detections = self.ner.detect(text)
        candidates: list[_Candidate] = []
        pending: list[SensitiveDetection] = []
        for detection in detections:
            if self._is_protected(detection.kind, detection.value) or self.ner.suggest_only or detection.confidence < self.ner.threshold:
                pending.append(detection)
            else:
                candidates.append(self._candidate_for_detection(detection))
        selected = self._select_candidates(candidates)
        selected_ranges = [(candidate.start, candidate.end) for candidate in selected]
        for detection in pending:
            if not any(not (detection.end <= start or detection.start >= end) for start, end in selected_ranges):
                self._record_suggestion(suggestion_from_detection(detection, self.placeholder_for(detection.placeholder_kind, detection.value)))
        for candidate in selected:
            self._register_auto_match(candidate.kind)
        return _apply_candidates(text, selected), selected

    def apply_with_edits(self, text: str) -> tuple[str, list[_Candidate]]:
        original_length = len(text)
        origin_map: list[tuple[int, ...]] = [(index,) for index in range(original_length)]
        edits: list[_Candidate] = []

        text, rule_edits = self._apply_rule_stage(text)
        for candidate in rule_edits:
            span = _candidate_original_span(candidate, origin_map)
            if span:
                edits.append(
                    _Candidate(span[0], span[1], candidate.replacement, "transformed", priority=1)
                )
        origin_map = _advance_origin_map(origin_map, rule_edits)

        if self.auto_pii:
            text, deterministic_edits = self._apply_deterministic_stage(text)
            for candidate in deterministic_edits:
                span = _candidate_original_span(candidate, origin_map)
                if span:
                    edits.append(
                        _Candidate(span[0], span[1], candidate.replacement, "transformed", priority=1)
                    )
            origin_map = _advance_origin_map(origin_map, deterministic_edits)

        text, ner_edits = self._apply_ner_stage(text)
        for candidate in ner_edits:
            span = _candidate_original_span(candidate, origin_map)
            if span:
                edits.append(
                    _Candidate(span[0], span[1], candidate.replacement, "transformed", priority=1)
                )
        return text, sorted(edits, key=lambda item: (item.start, item.end))

    def apply(self, text: str) -> str:
        return self.apply_with_edits(text)[0]

    def apply_markup(self, text: str, markup: str) -> str:
        projection = _markdown_or_html_projection(text, markup)
        if not projection.has_markup or not projection.visible:
            return self.apply(text)
        projected_result, edits = self.apply_with_edits(projection.visible)
        if projected_result == projection.visible:
            return self.apply(text)
        result = _replace_projected_characters(text, projection.positions, edits)
        # A second pass is intentional: the projection excludes HTML tags and
        # Markdown syntax, while relevant attributes/URLs remain in raw text.
        # Placeholders are protected, so this cannot double-replace an entity.
        return self.apply(result)

    def apply_fuzzy(self, text: str, rule_index: int, rule: Rule) -> str:
        """Apply one fuzzy rule for backwards-compatible programmatic callers."""

        selected = self._select_candidates(self._fuzzy_candidates(text, rule_index, rule))
        if selected:
            self.counts[rule_index] += len(selected)
        return _apply_candidates(text, selected)

    def replacement_for_box(self, rule_index: int, original: str) -> str:
        rule = self.rules[rule_index]
        self.counts[rule_index] += 1
        if rule.preserve_case:
            return adapt_case(original, rule.target)
        return rule.target

    def effective_rule_counts(self) -> list[int]:
        return [count + absorbed for count, absorbed in zip(self.counts, self.overlap_absorbed)]

    def report(self, safe: bool = True) -> list[dict[str, object]]:
        if not safe:
            return [
                {
                    "source": rule.source,
                    "target": rule.target,
                    "matches": self.counts[index],
                    "absorbed_matches": self.overlap_absorbed[index],
                    "smart": rule.smart,
                    "case_sensitive": rule.case_sensitive,
                    "fuzzy_threshold": rule.fuzzy_threshold,
                }
                for index, rule in enumerate(self.rules)
            ]
        return [
            {
                "rule_id": f"rule-{index + 1:04d}",
                "matches": self.counts[index],
                "absorbed_matches": self.overlap_absorbed[index],
                "status": "matched" if self.effective_rule_counts()[index] else "unmatched",
                "replacement": "configured",
            }
            for index in range(len(self.rules))
        ]

    def automatic_report(self) -> list[dict[str, object]]:
        return [
            {
                "kind": kind,
                "matches": count,
                "replacement": self.placeholder_kind_placeholder(kind),
                "status": "replaced",
            }
            for kind, count in sorted(self.auto_counts.items())
        ]

    def placeholder_kind_placeholder(self, kind: str) -> str:
        prefix = placeholder_kind_for(kind)
        return f"{prefix}-0001@example.com" if prefix == "EMAIL" else f"{prefix}-0001"


def escape_md_cell(value: object) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", r"\|").replace("\r\n", "\n").replace("\n", "<br>")


TABLE_CONTEXT_LABELS = {
    "address": "address",
    "account": "account",
    "bank account": "account",
    "city": "location",
    "cliente": "organization",
    "client": "organization",
    "company": "organization",
    "contact": "person",
    "contacto": "person",
    "country": "location",
    "cuenta": "account",
    "dirección": "address",
    "domicilio": "address",
    "email": "email",
    "empresa": "organization",
    "id": "id_or_ruc",
    "location": "location",
    "nombre": "person",
    "name": "person",
    "phone": "phone",
    "proveedor": "organization",
    "provider": "organization",
    "ruc": "id_or_ruc",
    "supplier": "organization",
    "tel": "phone",
    "teléfono": "phone",
    "telephone": "phone",
    "ubicación": "location",
}


def _table_context_kind(value: object) -> str | None:
    label = str(value or "").strip().rstrip(":").casefold()
    return TABLE_CONTEXT_LABELS.get(label)


TABLE_CONTEXT_CANONICAL_LABELS = {
    "address": "Dirección",
    "account": "Cuenta",
    "email": "Email",
    "id_or_ruc": "ID",
    "location": "Ubicación",
    "organization": "Empresa",
    "person": "Nombre",
    "phone": "Teléfono",
}


def _table_context_label(context_kind: str) -> str:
    return TABLE_CONTEXT_CANONICAL_LABELS.get(context_kind, context_kind)


def _apply_table_cell(value: object, engine: ReplacementEngine, context_kind: str | None = None) -> str:
    escaped = escape_md_cell(value)
    if not context_kind or not escaped:
        return engine.apply(escaped)

    label = _table_context_label(context_kind)
    prefix = f"{label}: "
    rendered = engine.apply_markup(prefix + escaped, "markdown")
    if rendered.casefold().startswith(prefix.casefold()):
        return rendered[len(prefix):]
    return engine.apply(escaped)


def markdown_table(rows: list[list[object]], engine: ReplacementEngine) -> str:
    if not rows:
        return ""

    width = max(len(row) for row in rows)
    normalized = [list(row) + [""] * (width - len(row)) for row in rows]
    header_contexts = {
        index: context
        for index, cell in enumerate(normalized[0])
        if (context := _table_context_kind(cell)) is not None
    }
    escaped: list[list[str]] = []
    for row_index, row in enumerate(normalized):
        processed: list[str] = []
        for column_index, cell in enumerate(row):
            cell_is_label = _table_context_kind(cell) is not None
            pair_context = _table_context_kind(row[column_index - 1]) if column_index and not cell_is_label else None
            context = pair_context or (header_contexts.get(column_index) if row_index and not cell_is_label else None)
            processed.append(_apply_table_cell(cell, engine, context))
        escaped.append(processed)

    header = escaped[0]
    body = escaped[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def trim_empty_edges(rows: list[list[object]]) -> list[list[object]]:
    while rows and all(cell in (None, "") for cell in rows[-1]):
        rows.pop()
    if not rows:
        return []

    max_width = max(len(row) for row in rows)
    while max_width > 0:
        if any(len(row) >= max_width and row[max_width - 1] not in (None, "") for row in rows):
            break
        max_width -= 1
    return [row[:max_width] for row in rows]


def text_file_to_markdown(path: Path, engine: ReplacementEngine) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() in {".md", ".markdown"}:
        return engine.apply_markup(text, "markdown")
    if path.suffix.lower() in {".html", ".htm"}:
        return engine.apply_markup(text, "html")
    return engine.apply(text)


def csv_to_markdown(path: Path, engine: ReplacementEngine, document_label: str = DEFAULT_DOCUMENT_LABEL) -> str:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        rows = list(csv.reader(handle, delimiter=delimiter))
    return f"# {document_label}\n\n" + markdown_table(rows, engine) + "\n"


def _xml_local_name(element: Any) -> str:
    tag = str(getattr(element, "tag", ""))
    return tag.rsplit("}", 1)[-1]


def _xml_attribute(element: Any, name: str) -> str:
    for key, value in getattr(element, "attrib", {}).items():
        if str(key).rsplit("}", 1)[-1] == name:
            return str(value)
    return ""


def _docx_run_property(run_element: Any, property_name: str) -> bool:
    for child in run_element:
        if _xml_local_name(child) != "rPr":
            continue
        for prop in child:
            if _xml_local_name(prop) == property_name:
                value = _xml_attribute(prop, "val").casefold()
                return value not in {"0", "false", "off", "none"}
    return False


def _docx_has_deleted_ancestor(element: Any) -> bool:
    parent = getattr(element, "getparent", lambda: None)()
    while parent is not None:
        if _xml_local_name(parent) in {"del", "moveFrom"}:
            return True
        parent = getattr(parent, "getparent", lambda: None)()
    return False


def _docx_run_for_node(element: Any) -> Any | None:
    parent = getattr(element, "getparent", lambda: None)()
    while parent is not None:
        if _xml_local_name(parent) == "r":
            return parent
        parent = getattr(parent, "getparent", lambda: None)()
    return None


def _docx_text_segments(paragraph: Any) -> list[_DocxTextSegment]:
    """Extract XML text nodes and their original run styles in order."""

    segments: list[_DocxTextSegment] = []
    offset = 0
    paragraph_element = getattr(paragraph, "_p", None)
    if paragraph_element is None:
        return segments
    for node in paragraph_element.iter():
        if _docx_has_deleted_ancestor(node):
            continue
        local = _xml_local_name(node)
        if local == "t":
            value = str(getattr(node, "text", "") or "")
        elif local == "tab":
            value = "\t"
        elif local in {"br", "cr"}:
            value = "\n"
        else:
            continue
        if not value:
            continue
        run = _docx_run_for_node(node)
        hyperlink = False
        hyperlink_url: str | None = None
        parent = getattr(node, "getparent", lambda: None)()
        while parent is not None:
            if _xml_local_name(parent) == "hyperlink":
                hyperlink = True
                relationship_id = _xml_attribute(parent, "id")
                part = getattr(paragraph, "part", None)
                relationships = getattr(part, "rels", {}) if part is not None else {}
                try:
                    relationship = relationships.get(relationship_id)
                    hyperlink_url = str(getattr(relationship, "target_ref", "") or "") or None
                except (AttributeError, KeyError, TypeError):
                    hyperlink_url = None
                break
            parent = getattr(parent, "getparent", lambda: None)()
        segments.append(
            _DocxTextSegment(
                text=value,
                start=offset,
                end=offset + len(value),
                bold=bool(run is not None and _docx_run_property(run, "b")),
                italic=bool(run is not None and _docx_run_property(run, "i")),
                underline=bool(run is not None and _docx_run_property(run, "u")),
                highlight=bool(run is not None and _docx_run_property(run, "highlight")),
                hyperlink=hyperlink,
                hyperlink_url=hyperlink_url,
            )
        )
        offset += len(value)
    return segments


def _wrap_docx_text(value: str, segment: _DocxTextSegment) -> str:
    if not value:
        return ""
    wrapped = value
    if segment.underline:
        wrapped = f"<u>{wrapped}</u>"
    if segment.highlight:
        wrapped = f"<mark>{wrapped}</mark>"
    if segment.italic:
        wrapped = f"_{wrapped}_"
    if segment.bold:
        wrapped = f"**{wrapped}**"
    if segment.hyperlink_url:
        wrapped = f"[{wrapped}]({segment.hyperlink_url})"
    return wrapped


def _render_docx_paragraph(paragraph: Any, engine: ReplacementEngine) -> str:
    segments = _docx_text_segments(paragraph)
    if not segments:
        fallback = str(getattr(paragraph, "text", "") or "")
        return engine.apply(fallback) if fallback else ""
    original = "".join(segment.text for segment in segments)
    transformed, edits = engine.apply_with_edits(original)
    if not edits:
        return "".join(_wrap_docx_text(segment.text, segment) for segment in segments)

    output: list[str] = []
    for segment in segments:
        fragment = original[segment.start:segment.end]
        for edit in sorted(edits, key=lambda item: item.start, reverse=True):
            if edit.end <= segment.start or edit.start >= segment.end:
                continue
            local_start = max(edit.start, segment.start) - segment.start
            local_end = min(edit.end, segment.end) - segment.start
            if edit.start >= segment.start:
                fragment = fragment[:local_start] + edit.replacement + fragment[local_end:]
            else:
                fragment = fragment[:0] + fragment[local_end:]
        output.append(_wrap_docx_text(fragment, segment))
    # A conservative guard for unusual diff layouts: never return the source
    # entity if a replacement was applied but could not be mapped to a run.
    rendered = "".join(output)
    if transformed and any(edit.replacement in rendered for edit in edits):
        return rendered
    return engine.apply(original)


def _docx_container_blocks(container: Any) -> Iterable[Any]:
    """Yield body/header/footer paragraphs and tables in document order."""

    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    element = getattr(container, "element", None)
    if element is None:
        element = getattr(container, "_element", None)
    if element is None:
        return
    for child in element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, container)
        elif isinstance(child, CT_Tbl):
            yield Table(child, container)


def _docx_table_markdown(table: Any, engine: ReplacementEngine) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        values: list[str] = []
        for cell in row.cells:
            paragraphs = [_render_docx_paragraph(paragraph, engine) for paragraph in cell.paragraphs]
            values.append("<br>".join(value.strip() for value in paragraphs if value.strip()))
        rows.append(values)
    if rows:
        header_contexts = {
            index: context
            for index, cell in enumerate(rows[0])
            if (context := _table_context_kind(cell)) is not None
        }
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                if _table_context_kind(value) is not None:
                    continue
                pair_context = _table_context_kind(row[column_index - 1]) if column_index else None
                context = pair_context or (header_contexts.get(column_index) if row_index else None)
                if context and value:
                    label = _table_context_label(context)
                    prefix = f"{label}: "
                    rendered = engine.apply_markup(prefix + value, "markdown")
                    if rendered.casefold().startswith(prefix.casefold()):
                        row[column_index] = rendered[len(prefix):]
    return _markdown_table_already_processed(rows)


def _markdown_table_already_processed(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [list(row) + [""] * (width - len(row)) for row in rows]
    lines = [
        "| " + " | ".join(escape_md_cell(cell) for cell in normalized[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(escape_md_cell(cell) for cell in row) + " |" for row in normalized[1:])
    return "\n".join(lines)


def _docx_auxiliary_text(path: Path, engine: ReplacementEngine) -> list[tuple[str, str]]:
    """Read comments/notes unavailable through python-docx's public API."""

    results: list[tuple[str, str]] = []
    members = {
        "word/comments.xml": "Comments",
        "word/footnotes.xml": "Footnotes",
        "word/endnotes.xml": "Endnotes",
    }
    try:
        with zipfile.ZipFile(path) as archive:
            for member, label in members.items():
                if member not in archive.namelist():
                    continue
                root = ET.fromstring(archive.read(member))
                values = [str(node.text or "") for node in root.iter() if _xml_local_name(node) == "t" and node.text]
                value = engine.apply(" ".join(values).strip())
                if value.strip():
                    results.append((label, value))
    except (OSError, KeyError, ET.ParseError, zipfile.BadZipFile):
        # Main document extraction remains useful even when an optional OOXML
        # part is malformed or unsupported.
        return results
    return results


def docx_to_markdown(path: Path, engine: ReplacementEngine, document_label: str = DEFAULT_DOCUMENT_LABEL) -> str:
    docx = require_import("docx", "python -m pip install python-docx")
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))

    def blocks(parent: Any):
        if parent is document:
            parent_elm = document.element.body
        elif isinstance(parent, Table):
            parent_elm = parent._tbl
        else:
            parent_elm = getattr(parent, "_element", None)
            if parent_elm is None:
                parent_elm = getattr(parent, "element", None)
        if parent_elm is None:
            return
        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    parts = [f"# {document_label}", ""]
    for block in blocks(document):
        if isinstance(block, Paragraph):
            text = _render_docx_paragraph(block, engine).strip()
            if text:
                parts.append(text)
                parts.append("")
        elif isinstance(block, Table):
            table_md = _docx_table_markdown(block, engine)
            if table_md:
                parts.append(table_md)
                parts.append("")

    seen_containers: set[int] = set()
    for section in document.sections:
        for role, container in (
            ("Header", section.header),
            ("Footer", section.footer),
        ):
            container_element = getattr(container, "_element", None)
            marker = id(container_element)
            if marker in seen_containers:
                continue
            seen_containers.add(marker)
            rendered: list[str] = []
            for block in blocks(container):
                if isinstance(block, Paragraph):
                    value = _render_docx_paragraph(block, engine).strip()
                    if value:
                        rendered.append(value)
                elif isinstance(block, Table):
                    value = _docx_table_markdown(block, engine)
                    if value:
                        rendered.append(value)
            if rendered:
                parts.extend([f"## {role}", "", "\n\n".join(rendered), ""])

    alternatives: list[str] = []
    for element in document.element.body.iter():
        if _xml_local_name(element) not in {"docPr", "cNvPr"}:
            continue
        for attribute in ("descr", "title"):
            value = _xml_attribute(element, attribute).strip()
            if value and value not in alternatives:
                alternatives.append(engine.apply(value))
    if alternatives:
        parts.extend(["## Text alternatives", "", *[f"- {value}" for value in alternatives], ""])

    for label, value in _docx_auxiliary_text(path, engine):
        parts.extend([f"## {label}", "", value, ""])

    return "\n".join(parts).strip() + "\n"


def xlsx_to_markdown(path: Path, engine: ReplacementEngine, document_label: str = DEFAULT_DOCUMENT_LABEL) -> str:
    openpyxl = require_import("openpyxl", "python -m pip install openpyxl")
    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts = [f"# {document_label}", ""]

    for sheet in workbook.worksheets:
        rows = trim_empty_edges([list(row) for row in sheet.iter_rows(values_only=True)])
        parts.append(f"## {engine.apply(sheet.title)}")
        parts.append("")
        if rows:
            parts.append(markdown_table(rows, engine))
        else:
            parts.append("_Empty sheet._")
        parts.append("")

    return "\n".join(parts).strip() + "\n"


def render_pdf_page_to_pil(page, dpi: int):
    fitz = require_import("fitz", "python -m pip install pymupdf")
    Image = require_import("PIL.Image", "python -m pip install pillow")
    scale = dpi / 72
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    return Image.open(io.BytesIO(pix.tobytes("png")))


def ocr_image_to_text(image, lang: str) -> str:
    pytesseract = require_import("pytesseract", "python -m pip install pytesseract pillow")
    return pytesseract.image_to_string(image, lang=lang)


def pdf_text_layer_text(page: Any) -> str:
    """Extract PDF text from ordered spans when a text layer exists."""

    try:
        data = page.get_text("dict")
        lines: list[str] = []
        for block in data.get("blocks", []) if isinstance(data, dict) else []:
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = [str(span.get("text", "")) for span in line.get("spans", []) if span.get("text")]
                if spans:
                    lines.append("".join(spans))
        if lines:
            return "\n".join(lines)
    except (AttributeError, KeyError, TypeError, ValueError):
        pass
    return str(page.get_text("text") or "")


def pdf_to_markdown(
    path: Path,
    engine: ReplacementEngine,
    ocr: bool,
    force_ocr: bool,
    lang: str,
    dpi: int,
    document_label: str = DEFAULT_DOCUMENT_LABEL,
) -> str:
    fitz = require_import("fitz", "python -m pip install pymupdf")
    doc = fitz.open(str(path))
    parts = [f"# {document_label}", ""]

    for index, page in enumerate(doc, start=1):
        used_ocr = False
        text = "" if force_ocr else pdf_text_layer_text(page).strip()
        if not text:
            if ocr or force_ocr:
                image = render_pdf_page_to_pil(page, dpi=dpi)
                text = ocr_image_to_text(image, lang=lang).strip()
                used_ocr = True
            else:
                text = "[No text layer found. Re-run with --ocr after installing Tesseract.]"

        parts.append(f"## Page {index}")
        if used_ocr:
            parts.append("")
            parts.append("<!-- OCR page: review accuracy before sharing. -->")
        parts.append("")
        parts.append(engine.apply(text))
        parts.append("")

    return "\n".join(parts).strip() + "\n"


def image_to_markdown(path: Path, engine: ReplacementEngine, lang: str, document_label: str = DEFAULT_DOCUMENT_LABEL) -> str:
    Image = require_import("PIL.Image", "python -m pip install pillow pytesseract")
    image = Image.open(str(path))
    text = ocr_image_to_text(image, lang=lang).strip()
    return f"# {document_label}\n\n<!-- OCR image: review accuracy before sharing. -->\n\n{engine.apply(text)}\n"


def to_markdown(
    path: Path,
    engine: ReplacementEngine,
    ocr: bool,
    force_ocr: bool,
    lang: str,
    dpi: int,
    document_label: str = DEFAULT_DOCUMENT_LABEL,
) -> str:
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        return text_file_to_markdown(path, engine)
    if ext in CSV_EXTS:
        return csv_to_markdown(path, engine, document_label=document_label)
    if ext in DOCX_EXTS:
        return docx_to_markdown(path, engine, document_label=document_label)
    if ext in XLSX_EXTS:
        return xlsx_to_markdown(path, engine, document_label=document_label)
    if ext in PDF_EXTS:
        return pdf_to_markdown(
            path,
            engine,
            ocr=ocr,
            force_ocr=force_ocr,
            lang=lang,
            dpi=dpi,
            document_label=document_label,
        )
    if ext in IMAGE_EXTS:
        return image_to_markdown(path, engine, lang=lang, document_label=document_label)
    die(f"unsupported input type '{ext}'")


def pdf_text_layer_boxes(page, dpi: int) -> list[WordBox]:
    scale = dpi / 72
    words = page.get_text("words")
    words.sort(key=lambda item: (item[5], item[6], item[7], item[1], item[0]))
    return [
        WordBox(
            text=str(word[4]),
            left=float(word[0]) * scale,
            top=float(word[1]) * scale,
            right=float(word[2]) * scale,
            bottom=float(word[3]) * scale,
            confidence=100.0,
        )
        for word in words
        if str(word[4]).strip()
    ]


def ocr_word_boxes(image, lang: str, min_conf: float) -> list[WordBox]:
    pytesseract = require_import("pytesseract", "python -m pip install pytesseract pillow")
    data = pytesseract.image_to_data(image, lang=lang, output_type=pytesseract.Output.DICT)
    boxes: list[WordBox] = []

    for index, text in enumerate(data.get("text", [])):
        text = str(text).strip()
        if not text:
            continue
        try:
            confidence = float(data["conf"][index])
        except (ValueError, KeyError):
            confidence = -1
        if confidence >= 0 and confidence < min_conf:
            continue
        left = float(data["left"][index])
        top = float(data["top"][index])
        width = float(data["width"][index])
        height = float(data["height"][index])
        boxes.append(WordBox(text=text, left=left, top=top, right=left + width, bottom=top + height, confidence=confidence))

    return boxes


def matching_box_groups(words: list[WordBox], rule: Rule) -> list[list[int]]:
    target = "".join(normalize_token(token, case_sensitive=rule.case_sensitive) for token in source_tokens(rule.source))
    if not target:
        return []

    normalized = [(index, normalize_token(word.text, case_sensitive=rule.case_sensitive)) for index, word in enumerate(words)]
    normalized = [(index, token) for index, token in normalized if token]
    matches: list[list[int]] = []
    used: set[int] = set()

    for start in range(len(normalized)):
        actual_indices: list[int] = []
        collected = ""
        for cursor in range(start, len(normalized)):
            word_index, token = normalized[cursor]
            if word_index in used:
                break
            collected += token
            actual_indices.append(word_index)
            is_match = collected == target
            if not is_match and len(target) >= 5 and rule.fuzzy_threshold > 0:
                is_match = similarity(collected, target) >= rule.fuzzy_threshold
            if is_match:
                matches.append(actual_indices)
                used.update(actual_indices)
                break
            if len(collected) > max(len(target) + 6, int(len(target) * 1.45)):
                break
            if rule.fuzzy_threshold <= 0 and not target.startswith(collected):
                break

    return matches


def parse_hex_color(value: str) -> tuple[int, int, int]:
    value = value.strip().lstrip("#")
    if len(value) != 6:
        die(f"invalid color '#{value}'. Use RRGGBB hex.")
    try:
        return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))
    except ValueError:
        die(f"invalid color '#{value}'. Use RRGGBB hex.")


def load_font(size: int):
    ImageFont = require_import("PIL.ImageFont", "python -m pip install pillow")
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def draw_fitted_text(draw, box: tuple[int, int, int, int], text: str, fill: tuple[int, int, int]) -> None:
    left, top, right, bottom = box
    max_width = max(1, right - left - 6)
    max_height = max(1, bottom - top - 4)
    size = max(8, int(max_height * 0.85))
    font = load_font(size)

    while size > 6:
        bbox = draw.textbbox((0, 0), text, font=font)
        if bbox[2] - bbox[0] <= max_width and bbox[3] - bbox[1] <= max_height:
            break
        size -= 1
        font = load_font(size)

    draw.text((left + 3, top + 2), text, fill=fill, font=font)


def parse_region(value: str) -> tuple[int, int, int, int, str]:
    raw_box, _, label = value.partition("=")
    parts = [part.strip() for part in raw_box.split(",")]
    if len(parts) != 4:
        die("invalid --redact-region. Use x,y,width,height or x,y,width,height=Label")
    try:
        x, y, width, height = [int(float(part)) for part in parts]
    except ValueError:
        die("invalid --redact-region numbers. Use x,y,width,height")
    if width <= 0 or height <= 0:
        die("invalid --redact-region size. width/height must be positive")
    return x, y, x + width, y + height, label.strip() or "REDACTED"


def apply_visual_redactions(
    image,
    engine: ReplacementEngine,
    redact_header: float,
    header_text: str,
    regions: list[str] | None,
    fill_color: tuple[int, int, int],
    text_color: tuple[int, int, int],
):
    if not redact_header and not regions:
        return image
    ImageDraw = require_import("PIL.ImageDraw", "python -m pip install pillow")
    image = image.convert("RGB")
    draw = ImageDraw.Draw(image)

    if redact_header:
        if redact_header <= 1:
            header_height = int(image.height * redact_header)
        else:
            header_height = int(redact_header)
        header_height = max(1, min(image.height, header_height))
        draw.rectangle((0, 0, image.width, header_height), fill=fill_color)
        if header_text:
            draw_fitted_text(draw, (8, 4, image.width - 8, header_height - 4), header_text, text_color)
        engine.visual_redactions += 1

    for region in regions or []:
        left, top, right, bottom, label = parse_region(region)
        left = max(0, min(image.width, left))
        right = max(0, min(image.width, right))
        top = max(0, min(image.height, top))
        bottom = max(0, min(image.height, bottom))
        draw.rectangle((left, top, right, bottom), fill=fill_color)
        if label:
            draw_fitted_text(draw, (left, top, right, bottom), label, text_color)
        engine.visual_redactions += 1

    return image


def redact_image_with_boxes(
    image,
    words: list[WordBox],
    engine: ReplacementEngine,
    fill_color: tuple[int, int, int],
    text_color: tuple[int, int, int],
):
    ImageDraw = require_import("PIL.ImageDraw", "python -m pip install pillow")
    image = image.convert("RGB")
    draw = ImageDraw.Draw(image)
    used_word_indices: set[int] = set()

    def draw_group(group: list[int], replacement: str) -> None:
        selected_words = [words[index] for index in group]
        left = int(max(0, min(word.left for word in selected_words) - 3))
        top = int(max(0, min(word.top for word in selected_words) - 2))
        right = int(min(image.width, max(word.right for word in selected_words) + 3))
        bottom = int(min(image.height, max(word.bottom for word in selected_words) + 2))
        draw.rectangle((left, top, right, bottom), fill=fill_color)
        draw_fitted_text(draw, (left, top, right, bottom), replacement, text_color)

    for rule_index, rule in enumerate(engine.rules):
        for group in matching_box_groups(words, rule):
            if any(index in used_word_indices for index in group):
                continue
            used_word_indices.update(group)
            selected = [words[index] for index in group]
            original = " ".join(word.text for word in selected)
            replacement = engine.replacement_for_box(rule_index, original)
            draw_group(group, replacement)

    # OCR and PDF text-layer words are also fed through the deterministic
    # classifier.  This keeps visual redaction aligned with text output and
    # never requires a remote OCR/NER service.
    if words:
        word_ranges: list[tuple[int, int]] = []
        cursor = 0
        for word in words:
            start = cursor
            cursor += len(word.text)
            word_ranges.append((start, cursor))
            cursor += 1
        joined = " ".join(word.text for word in words)
        detections = [detection for detection in detect_sensitive_content(joined) if detection.automatic]
        if engine.ner is not None:
            detections.extend(engine.ner.detect(joined))
        selected_detections: list[SensitiveDetection] = []
        for detection in sorted(
            detections,
            key=lambda item: (-(item.end - item.start), -DETECTION_PRIORITIES.get(item.kind, 0), item.start),
        ):
            if engine._is_protected(detection.kind, detection.value):
                continue
            if engine.ner is not None and detection.source == "spacy":
                if engine.ner.suggest_only or detection.confidence < engine.ner.threshold:
                    engine._record_suggestion(
                        suggestion_from_detection(detection, engine.placeholder_for(detection.placeholder_kind, detection.value))
                    )
                    continue
            group = [
                index
                for index, (start, end) in enumerate(word_ranges)
                if end > detection.start and start < detection.end
            ]
            if not group or any(index in used_word_indices for index in group):
                continue
            if any(
                not (detection.end <= previous.start or detection.start >= previous.end)
                for previous in selected_detections
            ):
                continue
            selected_detections.append(detection)
            used_word_indices.update(group)
            draw_group(group, engine.replacement_for_detection(detection))

    return image


def render_images(
    path: Path,
    output_dir: Path,
    engine: ReplacementEngine,
    ocr: bool,
    force_ocr: bool,
    lang: str,
    dpi: int,
    min_conf: float,
    image_fill: str,
    image_text: str,
    redact_header: float,
    header_text: str,
    redact_regions: list[str] | None,
    document_label: str = DEFAULT_DOCUMENT_LABEL,
    progress: bool = False,
) -> list[Path]:
    ext = path.suffix.lower()
    fill_color = parse_hex_color(image_fill)
    text_color = parse_hex_color(image_text)
    output_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    engine.visual_review_required = True
    output_label = safe_label_filename(document_label)

    if ext in PDF_EXTS:
        fitz = require_import("fitz", "python -m pip install pymupdf")
        doc = fitz.open(str(path))
        for index, page in enumerate(doc, start=1):
            image = render_pdf_page_to_pil(page, dpi=dpi)
            boxes = [] if force_ocr else pdf_text_layer_boxes(page, dpi=dpi)
            if not boxes:
                if not ocr and not force_ocr:
                    die("page has no text boxes. Re-run with --ocr for scanned PDFs.")
                boxes = ocr_word_boxes(image, lang=lang, min_conf=min_conf)
            redacted = redact_image_with_boxes(image, boxes, engine, fill_color, text_color)
            redacted = apply_visual_redactions(redacted, engine, redact_header, header_text, redact_regions, fill_color, text_color)
            out_path = output_dir / f"{output_label}-page-{index:04d}.png"
            redacted.save(str(out_path))
            written.append(out_path)
            if progress:
                print(f"  page {index}: wrote {out_path.name}", file=sys.stderr)
        return written

    if ext in IMAGE_EXTS:
        Image = require_import("PIL.Image", "python -m pip install pillow pytesseract")
        image = Image.open(str(path))
        boxes = ocr_word_boxes(image, lang=lang, min_conf=min_conf)
        redacted = redact_image_with_boxes(image, boxes, engine, fill_color, text_color)
        redacted = apply_visual_redactions(redacted, engine, redact_header, header_text, redact_regions, fill_color, text_color)
        out_path = output_dir / f"{output_label}-image.png"
        redacted.save(str(out_path))
        written.append(out_path)
        if progress:
            print(f"  image: wrote {out_path.name}", file=sys.stderr)
        return written

    die("--format images only supports PDF or image inputs")


def parse_rule_text(raw: str, case_sensitive: bool, smart: bool, preserve_case: bool, fuzzy_threshold: float) -> Rule:
    if "=" not in raw:
        die("invalid --replace rule; use source=target")
    source, target = raw.split("=", 1)
    source = source.strip()
    target = target.strip()
    if not source:
        die("replacement source cannot be empty")
    return Rule(source=source, target=target, case_sensitive=case_sensitive, smart=smart, preserve_case=preserve_case, fuzzy_threshold=fuzzy_threshold)


def load_rules(args: argparse.Namespace) -> list[Rule]:
    rules: list[Rule] = []

    for raw in args.replace or []:
        rules.append(
            parse_rule_text(
                raw,
                case_sensitive=args.case_sensitive,
                smart=not args.exact,
                preserve_case=not args.no_preserve_case,
                fuzzy_threshold=0.0 if args.no_fuzzy else args.fuzzy_threshold,
            )
        )

    if args.rules:
        data = json.loads(Path(args.rules).read_text(encoding="utf-8"))
        items = data.items() if isinstance(data, dict) else data
        for item in items:
            if isinstance(item, tuple):
                source, target = item
                item_options = {}
            else:
                source = item.get("source", item.get("from"))
                target = item.get("target", item.get("to"))
                item_options = item
            if not source:
                die("every JSON rule needs 'source'/'from'")
            if target is None:
                die("every JSON rule needs 'target'/'to'")
            rules.append(
                Rule(
                    source=str(source),
                    target=str(target),
                    case_sensitive=bool(item_options.get("case_sensitive", args.case_sensitive)),
                    smart=bool(item_options.get("smart", not args.exact)),
                    preserve_case=bool(item_options.get("preserve_case", not args.no_preserve_case)),
                    fuzzy_threshold=float(item_options.get("fuzzy_threshold", 0.0 if args.no_fuzzy else args.fuzzy_threshold)),
                )
            )

    return rules


def resolve_text_output(input_path: Path, output: str | None, suffix: str) -> Path:
    if output is None:
        return input_path.with_name(f"document.nubli{suffix}")
    out = Path(output)
    if out.suffix:
        return out
    return out / f"document.nubli{suffix}"


def resolve_image_output(input_path: Path, output: str | None) -> Path:
    if output is None:
        return input_path.with_name("document.nubli-images")
    return Path(output)


def normalize_document_label(value: str | None) -> str:
    label = (value or DEFAULT_DOCUMENT_LABEL).strip().replace("\r", " ").replace("\n", " ")
    return label or DEFAULT_DOCUMENT_LABEL


def safe_label_filename(value: str) -> str:
    label = re.sub(r"[^A-Za-z0-9._-]+", "-", normalize_document_label(value)).strip(".-")
    return label or "DOCUMENT-001"


NUBLI_BANNER = """
┌─ Nubli ───────────────────────────────┐
│ local-first document anonymizer        │
└───────────────────────────────────────┘
"""

NUBLI_PET = "  /ᐠ｡ꞈ｡ᐟ\\  redact before you chat"


def print_install_help() -> None:
    print(NUBLI_BANNER.rstrip())
    print("Install / download Nubli CLI\n")
    print("If you already have this repo/folder:")
    print("  cd nubli")
    print("  ./install.sh")
    print("\nAfter install:")
    print("  nubli")
    print("  nubli --help")
    print("\nFor a future GitHub release:")
    print("  git clone <repo-url> nubli")
    print("  cd nubli")
    print("  ./install.sh")
    print("\nHide installer mascot:")
    print("  NUBLI_NO_PET=1 ./install.sh")
    print("  ./install.sh --no-pet")


def ui(args: argparse.Namespace, message: str = "") -> None:
    if not getattr(args, "quiet", False):
        print(message, file=sys.stderr)


def show_run_header(args: argparse.Namespace, input_path: Path, menu_opened: bool) -> None:
    if args.quiet:
        return
    if not menu_opened:
        print(NUBLI_BANNER.rstrip(), file=sys.stderr)
        if not args.no_pet:
            print(NUBLI_PET, file=sys.stderr)
    print("nubli: private docs stay local", file=sys.stderr)
    print("[1/5] input   : local document (name suppressed)", file=sys.stderr)
    print(f"      label   : {args.document_label}", file=sys.stderr)
    print(f"      format  : {args.format}", file=sys.stderr)
    if args.output:
        print(f"      output  : {args.output}", file=sys.stderr)
    if args.ocr or args.force_ocr:
        print(f"      OCR     : {'force' if args.force_ocr else 'fallback'} ({args.ocr_lang})", file=sys.stderr)


def show_rule_summary(args: argparse.Namespace, rules: list[Rule]) -> None:
    if args.quiet:
        return
    print(f"[2/5] rules   : {len(rules)} replacement rule(s)", file=sys.stderr)
    if args.rules:
        print("      file    : configured local rules (path suppressed)", file=sys.stderr)
    if rules:
        print(f"      fuzzy   : {'off' if args.no_fuzzy else args.fuzzy_threshold}", file=sys.stderr)
    if args.redact_header:
        print(f"      header  : redact {args.redact_header}", file=sys.stderr)
    for region in args.redact_region or []:
        print(f"      region  : {region}", file=sys.stderr)


def show_match_summary(args: argparse.Namespace, engine: ReplacementEngine) -> None:
    if args.quiet:
        return
    print("[4/5] matches :", file=sys.stderr)
    if not engine.rules:
        print("      no explicit text rules; deterministic detectors active", file=sys.stderr)
    for index, count in enumerate(engine.counts):
        effective = engine.effective_rule_counts()[index]
        print(
            f"      rule-{index + 1:04d}: {count} direct match(es), "
            f"{engine.overlap_absorbed[index]} overlap(s) absorbed (effective={effective})",
            file=sys.stderr,
        )
    for kind, count in sorted(engine.auto_counts.items()):
        print(f"      {kind}: {count} automatic placeholder match(es)", file=sys.stderr)
    if engine.visual_redactions:
        print(f"      visual redactions: {engine.visual_redactions}", file=sys.stderr)


def show_done(args: argparse.Namespace, paths: list[Path], report_path: Path | None = None) -> None:
    if args.quiet:
        return
    if getattr(args, "format", "markdown") == "images":
        print("[5/5] done    : visual output ready; human review required", file=sys.stderr)
    else:
        print("[5/5] done    : anonymized output ready", file=sys.stderr)
    for path in paths:
        print(f"      output  : {path}", file=sys.stderr)
    if report_path:
        print(f"      report  : {report_path}", file=sys.stderr)


def interactive_menu() -> list[str]:
    print(NUBLI_BANNER.rstrip())
    print(NUBLI_PET)
    print("Private docs stay local. Choose an action:\n")
    print("  1) Anonymize a file")
    print("  2) Show install/download help")
    print("  3) Show quick examples")
    print("  4) Exit")
    choice = input("\nSelect [1-4] (default: 1): ").strip() or "1"

    if choice == "2":
        print()
        print_install_help()
        raise SystemExit(0)
    if choice == "3":
        print("\nQuick examples:\n")
        print('  nubli document.pdf --replace "Acme S A=Demo Company" -o out.md')
        print('  nubli document.pdf --format images --redact-header 0.12 -o out-pages')
        print('  nubli document.docx --rules rules.json --strict-pii -o out.md')
        raise SystemExit(0)
    if choice == "4":
        raise SystemExit(0)

    print("\nStep 1 — Input")
    input_path = input("File to anonymize: ").strip().strip('"')
    if not input_path:
        die("input file is required")

    print("\nStep 2 — Output format")
    print("  1) markdown  (best for AI/chat)")
    print("  2) text")
    print("  3) images    (best for PDFs with logos/headers)")
    fmt_choice = input("Select [1-3] (default: 1): ").strip() or "1"
    fmt = {"1": "markdown", "2": "text", "3": "images"}.get(fmt_choice, fmt_choice.lower())
    if fmt not in {"markdown", "text", "images"}:
        fmt = "markdown"

    print("\nStep 3 — Output path")
    output_path = input("Output file/folder (blank = auto): ").strip().strip('"')

    argv = [input_path, "--format", fmt]
    if output_path:
        argv.extend(["-o", output_path])

    print("\nStep 4 — Replacement rules")
    rules_path = input("Rules JSON path (blank = type rules manually): ").strip().strip('"')
    if rules_path:
        argv.extend(["--rules", rules_path])

    print("Add rules like: Acme S A=Demo Company")
    print("Leave blank when done.")
    while True:
        rule = input("Rule (blank to finish): ").strip()
        if not rule:
            break
        argv.extend(["--replace", rule])

    print("\nStep 5 — Safety")
    if prompt_yes_no("Require every rule to match?", default=True):
        argv.append("--require-all-rules")
    if prompt_yes_no("Block output if possible emails/phones/names remain?", default=True):
        argv.append("--strict-pii")

    if fmt == "images":
        print("\nStep 6 — PDF/image visual options")
        if prompt_yes_no("Use OCR fallback for scanned/image-only pages?", default=True):
            argv.append("--ocr")
        if prompt_yes_no("Redact top header area too? Useful for logos/company names.", default=False):
            pct = input("Header height percent (default 0.12): ").strip() or "0.12"
            argv.extend(["--redact-header", pct])
            label = input("Header replacement text (default REDACTED): ").strip()
            if label:
                argv.extend(["--header-text", label])
        while prompt_yes_no("Add a manual logo/region redaction?", default=False):
            region = input("Region x,y,width,height=Label: ").strip()
            if region:
                argv.extend(["--redact-region", region])

    print("\nRunning Nubli...\n")

    return argv


def suggestion_report_item(suggestion: Suggestion, safe: bool = True) -> dict[str, object]:
    if not safe:
        return suggestion.__dict__.copy()
    return {
        "kind": suggestion.kind,
        "matches": 1,
        "replacement": suggestion.target,
        "status": "suggestion",
        "confidence": round(suggestion.confidence, 4),
        "source": suggestion.source,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="nubli",
        description="Local-first anonymizer for PDF, image, DOCX, XLSX, CSV and Markdown files.",
    )
    parser.add_argument("input", nargs="?", help="Input file path. Omit to open the interactive menu.")
    parser.add_argument("--help-install", action="store_true", help="Show CLI install/download instructions")
    parser.add_argument("-o", "--output", help="Output file or folder")
    parser.add_argument("--format", choices=["markdown", "text", "images"], default="markdown", help="Output format")
    parser.add_argument("--replace", action="append", help='Replacement rule: "Private Name=Public Name"')
    parser.add_argument("--rules", help="JSON rules file. Accepts object or list of {from/source,to/target}.")
    parser.add_argument("--ocr", action="store_true", help="Use local OCR fallback for scanned PDFs/images")
    parser.add_argument("--force-ocr", action="store_true", help="Ignore PDF text layer and OCR every PDF page")
    parser.add_argument("--ocr-lang", default="eng+spa", help="Tesseract languages, e.g. eng, spa, eng+spa")
    parser.add_argument("--dpi", type=int, default=220, help="PDF render DPI for OCR/image output")
    parser.add_argument("--min-ocr-conf", type=float, default=45, help="Minimum OCR word confidence for image redaction")
    parser.add_argument("--case-sensitive", action="store_true", help="Make CLI --replace rules case-sensitive")
    parser.add_argument("--exact", action="store_true", help="Use literal matching instead of smart punctuation/space matching")
    parser.add_argument("--fuzzy-threshold", type=float, default=0.88, help="Fuzzy typo-match threshold from 0.0 to 1.0")
    parser.add_argument("--no-fuzzy", action="store_true", help="Disable fuzzy typo matching")
    parser.add_argument("--no-preserve-case", action="store_true", help="Always use target exactly as typed")
    parser.add_argument("--report", help="Write JSON replacement report")
    parser.add_argument("--safe-report", action="store_true", help="Keep the JSON report free of original values (default)")
    parser.add_argument("--unsafe-report", action="store_true", help="Include original rule values in the report (explicitly unsafe)")
    parser.add_argument("--document-label", default=DEFAULT_DOCUMENT_LABEL, help="Neutral title/label used in generated output")
    parser.add_argument("--enable-ner", action="store_true", help="Enable optional local spaCy NER; never downloads a model")
    parser.add_argument("--ner-model", help="Local spaCy model directory; no remote paths or downloads")
    parser.add_argument("--ner-language", default="es", help="NER language used for the default installed model")
    parser.add_argument("--ner-threshold", type=float, default=0.85, help="Minimum local NER confidence for automatic replacement")
    parser.add_argument("--suggest-only-ner", action="store_true", help="Report local NER entities as suggestions without replacing them")
    parser.add_argument("--quiet", action="store_true", help="Print only output paths/errors; hides banner and progress UI")
    parser.add_argument("--no-pet", action="store_true", help="Hide the Nubli cat mascot but keep progress UI")
    parser.add_argument("--require-all-rules", action="store_true", help="Fail/ask if any replacement rule had zero matches")
    parser.add_argument("--strict-pii", action="store_true", help="Fail/ask if output still looks like it contains emails, phones, IDs or names")
    parser.add_argument("--no-suggest-pii", dest="suggest_pii", action="store_false", help="Disable automatic sensitive-data suggestions")
    parser.set_defaults(suggest_pii=True, interactive=True)
    parser.add_argument("--no-interactive", dest="interactive", action="store_false", help="Never ask questions; fail closed instead")
    parser.add_argument(
        "--allow-zero-replacements",
        action="store_true",
        help="Allow output even when no replacement rule matched. By default Nubli fails closed to avoid sharing un-anonymized files.",
    )
    parser.add_argument("--image-fill", default="#ffffff", help="Redaction fill color for image output")
    parser.add_argument("--image-text", default="#111111", help="Replacement text color for image output")
    parser.add_argument("--redact-header", type=float, default=0.0, help="For image output, redact the top header by percent (0.12) or pixels (160)")
    parser.add_argument("--header-text", default="REDACTED", help="Text to write inside --redact-header area")
    parser.add_argument("--redact-region", action="append", help="For image output, redact x,y,width,height or x,y,width,height=Label. Repeatable for logos/seals.")
    return parser


def main(argv: list[str] | None = None) -> int:
    raw_argv = sys.argv[1:] if argv is None else argv
    menu_opened = False
    if not raw_argv:
        menu_opened = True
        raw_argv = interactive_menu()

    args = build_parser().parse_args(raw_argv)
    if args.help_install:
        print_install_help()
        return 0

    if not args.input:
        menu_opened = True
        raw_argv = interactive_menu()
        args = build_parser().parse_args(raw_argv)
        if args.help_install:
            print_install_help()
            return 0

    input_path = Path(args.input).expanduser().resolve()
    if not input_path.exists():
        die(f"input file does not exist: {input_path}")

    if not 0 <= args.fuzzy_threshold <= 1:
        die("--fuzzy-threshold must be between 0.0 and 1.0")
    if not 0 <= args.ner_threshold <= 1:
        die("--ner-threshold must be between 0.0 and 1.0")
    if args.safe_report and args.unsafe_report:
        die("--safe-report and --unsafe-report cannot be used together")
    args.document_label = normalize_document_label(args.document_label)
    if args.unsafe_report:
        print("WARNING: --unsafe-report may contain original sensitive values.", file=sys.stderr)

    show_run_header(args, input_path, menu_opened)
    rules = load_rules(args)
    show_rule_summary(args, rules)

    ner = LocalNER(
        enabled=bool(args.enable_ner or args.ner_model or args.suggest_only_ner),
        model_path=args.ner_model,
        language=args.ner_language,
        threshold=args.ner_threshold,
        suggest_only=args.suggest_only_ner,
    )
    engine = ReplacementEngine(rules, ner=ner)
    suggestions: list[Suggestion] = []
    output_paths: list[Path] = []

    if args.format in {"markdown", "text"}:
        suffix = ".md" if args.format == "markdown" else ".txt"
        out_path = resolve_text_output(input_path, args.output, suffix).resolve()
        out_path.parent.mkdir(parents=True, exist_ok=True)
        ui(args, "[3/5] process : extracting text and applying replacements")
        content = to_markdown(
            input_path,
            engine,
            ocr=args.ocr,
            force_ocr=args.force_ocr,
            lang=args.ocr_lang,
            dpi=args.dpi,
            document_label=args.document_label,
        )
        suggestions = merge_suggestions(
            engine.suggestions,
            suggest_sensitive_content(content) if args.suggest_pii else [],
        )
        show_match_summary(args, engine)
        enforce_safety(args, engine, suggestions)
        out_path.write_text(content, encoding="utf-8")
        output_paths.append(out_path)
        if args.quiet:
            print(out_path)
    else:
        output_dir = resolve_image_output(input_path, args.output).resolve()
        ui(args, "[3/5] process : rendering pages and applying visual redactions")
        written = render_images(
            input_path,
            output_dir,
            engine,
            ocr=args.ocr,
            force_ocr=args.force_ocr,
            lang=args.ocr_lang,
            dpi=args.dpi,
            min_conf=args.min_ocr_conf,
            image_fill=args.image_fill,
            image_text=args.image_text,
            redact_header=args.redact_header,
            header_text=args.header_text,
            redact_regions=args.redact_region,
            document_label=args.document_label,
            progress=not args.quiet,
        )
        suggestions = merge_suggestions(engine.suggestions)
        show_match_summary(args, engine)
        enforce_safety(args, engine, suggestions, written=written)
        for path in written:
            output_paths.append(path)
            if args.quiet:
                print(path)

    report_path: Path | None = None
    if args.report:
        report_path = Path(args.report).expanduser().resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        safe_report = not args.unsafe_report
        report = {
            "document_label": args.document_label,
            "rules": engine.report(safe=safe_report),
            "total_rule_matches": sum(engine.counts),
            "automatic": engine.automatic_report(),
            "overlap_events": engine.overlap_events,
            "visual_redactions": engine.visual_redactions,
            "visual_review_required": engine.visual_review_required,
            "ner": engine.ner.report() if engine.ner is not None else {"enabled": False},
            "suggestions": [suggestion_report_item(suggestion, safe=safe_report) for suggestion in suggestions],
        }
        report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        if args.quiet:
            print(report_path)

    show_done(args, output_paths, report_path)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
