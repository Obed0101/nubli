from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import nubli  # noqa: E402


class SemanticAnonymizationTests(unittest.TestCase):
    def test_synthetic_categories_and_precedence(self) -> None:
        source = (
            "(Ana Pérez) [Compañía Demo S.A.] {Dirección: Calle Falsa 123} "
            "“Cuenta bancaria: 0000000001” Teléfono: 0000-0000 "
            "(+507) 0000-0000 persona@example.test ID: 12-345-6789 "
            "RUC: 12345678-9 Fecha: 01/02/2024 Ciudad: Panamá "
            "Valor: secreto sintético"
        )
        detections = nubli.detect_sensitive_content(source)
        kinds = {d.kind for d in detections}
        self.assertTrue({"person", "organization", "address", "account", "phone", "email", "id_or_ruc", "date", "location", "value"} <= kinds)

        engine = nubli.ReplacementEngine([])
        result = engine.apply(source)
        for value in (
            "Ana Pérez",
            "Compañía Demo",
            "Calle Falsa 123",
            "0000000001",
            "0000-0000",
            "(+507) 0000-0000",
            "persona@example.test",
            "12-345-6789",
            "12345678-9",
            "01/02/2024",
            "Panamá",
            "secreto sintético",
        ):
            self.assertNotIn(value, result)
        for marker in ("PERSON-0001", "COMPANY-0001", "ADDRESS-0001", "ACCOUNT-0001", "PHONE-0001", "EMAIL-0001@example.com", "ID-0001", "DATE-0001", "LOCATION-0001", "VALUE-0001"):
            self.assertIn(marker, result)

        self.assertIn("<PERSON-0001>", nubli.ReplacementEngine([]).apply("<Ana Pérez>"))
        self.assertIn("ACCOUNT-0001", nubli.ReplacementEngine([]).apply("SWIFT: BOFAUS3N IBAN: ES9121000418450200051332"))
        self.assertIn("(", result)
        self.assertIn("]", result)
        self.assertIn("{Dirección:", result)
        self.assertIn("“Cuenta bancaria:", result)

    def test_country_phone_prefix_does_not_shadow_hyphenated_id(self) -> None:
        detections = nubli.detect_sensitive_content("Teléfono: 507-6000-0000 ID: 12-345-6789")
        self.assertIn(("phone", "507-6000-0000"), [(d.kind, d.value) for d in detections])
        self.assertIn(("id_or_ruc", "12-345-6789"), [(d.kind, d.value) for d in detections])

    def test_context_boundaries_do_not_consume_adjacent_fields(self) -> None:
        source = "Empresa: Organización Sintética proveedor: Organización Sintética Dirección: Calle Falsa 123 [Domicilio: Avenida Sintética 45]"
        detections = nubli.detect_sensitive_content(source)
        values = [(d.kind, d.value) for d in detections]
        self.assertEqual(values.count(("organization", "Organización Sintética")), 2)
        self.assertIn(("address", "Calle Falsa 123"), values)
        self.assertIn(("address", "Avenida Sintética 45"), values)
        self.assertEqual(nubli.ReplacementEngine([]).apply("Nombre Persona: proveedor"), "Nombre Persona: proveedor")

    def test_names_are_conservative_and_placeholder_mapping_is_stable(self) -> None:
        source = "Heading Texto Normal [Ana Pérez] Ana Pérez Company Demo Company"
        engine = nubli.ReplacementEngine([])
        result = engine.apply(source)
        self.assertEqual(result.count("PERSON-0001"), 2)
        self.assertIn("Texto Normal", result)
        self.assertIn("Demo Company", result)
        self.assertNotIn("Ana Pérez", result)

    def test_markdown_and_html_structure(self) -> None:
        markdown = "# Heading\n\n**Ana Pérez** _Compañía Demo S.A._ [Calle Falsa 123](https://example.test) `persona@example.test`"
        html = '<h1>Heading</h1><p><strong>Ana</strong> <em>Pérez</em></p><a href="mailto:persona@example.test">Compañía Demo S.A.</a><img alt="Ana Pérez" title="Calle Falsa 123">'
        markdown_result = nubli.ReplacementEngine([]).apply_markup(markdown, "markdown")
        html_result = nubli.ReplacementEngine([]).apply_markup(html, "html")
        for result in (markdown_result, html_result):
            self.assertNotIn("Ana Pérez", result)
            self.assertNotIn("persona@example.test", result)
            self.assertNotIn("Compañía Demo", result)
            self.assertNotIn("Calle Falsa 123", result)
        self.assertIn("**PERSON-0001**", markdown_result)
        self.assertIn("<strong>", html_result)
        self.assertIn("href=\"mailto:EMAIL-0001@example.com\"", html_result)
        self.assertIn("alt=\"PERSON-0001\"", html_result)

    def test_english_context_labels_work_without_ner(self) -> None:
        source = "Name: Jane Doe Company: Example Corporation Address: 1 Main Street Phone: 0000-0000 Email: jane@example.test ID: AB-12345 Account: 0000000001 Location: London"
        result = nubli.ReplacementEngine([]).apply(source)
        self.assertIn("PERSON-0001", result)
        self.assertIn("COMPANY-0001", result)
        self.assertIn("ADDRESS-0001", result)
        self.assertIn("PHONE-0001", result)
        self.assertIn("EMAIL-0001@example.com", result)
        self.assertIn("ID-0001", result)
        self.assertIn("ACCOUNT-0001", result)
        self.assertIn("LOCATION-0001", result)
        for value in ("Jane Doe", "Example Corporation", "1 Main Street", "jane@example.test", "London"):
            self.assertNotIn(value, result)

    def test_report_is_safe_and_input_hash_is_unchanged(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            source_path = directory_path / "synthetic.md"
            output_path = directory_path / "output.md"
            report_path = directory_path / "report.json"
            source_path.write_text("Ana Pérez; persona@example.test; Empresa: Organización Sintética\n", encoding="utf-8")
            before = hashlib.sha256(source_path.read_bytes()).hexdigest()
            completed = subprocess.run(
                [
                    sys.executable,
                    str(ROOT / "nubli.py"),
                    str(source_path),
                    "--format",
                    "markdown",
                    "--report",
                    str(report_path),
                    "--no-interactive",
                    "--quiet",
                    "-o",
                    str(output_path),
                ],
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertEqual(before, hashlib.sha256(source_path.read_bytes()).hexdigest())
            report = json.loads(report_path.read_text(encoding="utf-8"))
            report_text = json.dumps(report, ensure_ascii=False)
            self.assertNotIn("Ana Pérez", report_text)
            self.assertNotIn("persona@example.test", report_text)
            self.assertNotIn("Organización Sintética", report_text)
            self.assertEqual(report["document_label"], "DOCUMENT-001")
            self.assertNotIn("synthetic.md", output_path.read_text(encoding="utf-8"))

    def test_fail_closed_when_nothing_matches(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            source_path = directory_path / "plain.md"
            output_path = directory_path / "output.md"
            source_path.write_text("Texto público sin datos detectables.\n", encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, str(ROOT / "nubli.py"), str(source_path), "--no-interactive", "--quiet", "-o", str(output_path)],
                capture_output=True,
                text=True,
            )
            self.assertNotEqual(completed.returncode, 0)
            self.assertFalse(output_path.exists())

    def test_ner_remote_paths_are_rejected_without_network(self) -> None:
        ner = nubli.LocalNER(enabled=True, model_path="https://example.invalid/model")
        self.assertIsNone(ner.load())
        self.assertEqual(ner.error, "remote model paths are disabled")
        self.assertTrue(ner.report()["fallback"])

    @unittest.skipUnless(__import__("importlib").util.find_spec("openpyxl"), "optional openpyxl dependency is not installed")
    def test_xlsx_tables_are_redacted(self) -> None:
        import openpyxl

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Synthetic"
            sheet.append(["Proveedor", "Compañía Demo"])
            workbook.save(path)
            output = nubli.xlsx_to_markdown(path, nubli.ReplacementEngine([]))
            self.assertIn("COMPANY-0001", output)
            self.assertNotIn("Compañía Demo", output)

    @unittest.skipUnless(__import__("importlib").util.find_spec("docx"), "optional python-docx dependency is not installed")
    def test_docx_runs_headers_and_footers_are_redacted(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.docx"
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("Ana").bold = True
            paragraph.add_run(" Pérez").italic = True
            document.sections[0].header.paragraphs[0].add_run("Juan Pérez")
            document.sections[0].footer.paragraphs[0].add_run("persona@example.test")
            document.save(path)
            output = nubli.docx_to_markdown(path, nubli.ReplacementEngine([]))
            self.assertNotIn("Ana Pérez", output)
            self.assertNotIn("Juan Pérez", output)
            self.assertNotIn("persona@example.test", output)
            self.assertIn("PERSON-0001", output)
            self.assertIn("EMAIL-0001@example.com", output)


if __name__ == "__main__":
    unittest.main()
